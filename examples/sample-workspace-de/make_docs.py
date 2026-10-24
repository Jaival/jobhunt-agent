"""Write the German sample master CV and cover letter (cv/Jonas_Becker_Lebenslauf.docx, cv/Jonas_Becker_Anschreiben.docx).

Jonas Becker is fictional: a mechanical engineer looking for junior simulation roles in Bayern. His CV is plain
paragraphs with German headings: project lines are "Title<TAB>dates" followed by Word list bullets, and it ends with
a "City, date" line. His letter has a multi-line address block and no empty line before the closing. The master
files say "Jonas Michael Becker"; the applications say "Jonas Becker" (config.yaml name_in_master).

Usage: python make_docs.py [workspace folder]   (default: the folder this file is in)
"""
import sys
from pathlib import Path

import docx
from docx.shared import Pt

CV = [
    ("Jonas Michael Becker", "name"),
    ("Maschinenbauingenieur | FEM-Simulation, CAD & Konstruktion", None),
    ("Regensburg | jonas.becker@example.com | +49 160 00000000", None),
    ("", None),
    ("PROFIL", "heading"),
    ("Maschinenbauingenieur (B.Eng.) mit Erfahrung in FEM-Berechnung und Konstruktion von Baugruppen.", None),
    ("", None),
    ("TECHNISCHE KOMPETENZEN", "heading"),
    ("Simulation: ANSYS Mechanical, Abaqus", None),
    ("CAD: SolidWorks, CATIA V5", None),
    ("Programmierung: Python, MATLAB", None),
    ("Sprachen: Deutsch (Muttersprache), Englisch (C1)", None),
    ("", None),
    ("BERUFSERFAHRUNG", "heading"),
    ("Werkstudent Berechnung, Muster Automotive GmbH, Regensburg\t10/2023 – 03/2025", "bold"),
    ("Festigkeitsnachweise für Halterungen mit ANSYS erstellt.", "bullet"),
    ("Berechnungsberichte für die Konstruktion zusammengefasst.", "bullet"),
    ("", None),
    ("PROJEKTE", "heading"),
    ("FEM-Analyse Radträger, Bachelorarbeit\t03/2025 – 08/2025", "bold"),
    ("Topologieoptimierung eines Radträgers, 18 % weniger Masse.", "bullet"),
    ("Validierung mit Dehnungsmessstreifen am Prüfstand.", "bullet"),
    ("Formula Student Fahrwerk\t2022 – 2024", "bold"),
    ("Konstruktion der Querlenker in CATIA V5.", "bullet"),
    ("Python-Skript für Kinematikauswertung\t2023", "bold"),
    ("Auswertung von Messdaten aus der Fahrwerkskinematik.", "bullet"),
    ("", None),
    ("AUSBILDUNG", "heading"),
    ("B.Eng. Maschinenbau, OTH Regensburg\t10/2021 – 08/2025", "bold"),
    ("", None),
    ("Regensburg, 01.10.2026", None),
]

LETTER = [
    "Jonas Michael Becker",
    "Musterweg 3, 93047 Regensburg",
    "jonas.becker@example.com | +49 160 00000000",
    "",
    "Beispiel Werke AG",
    "Hauptstraße 1",
    "80331 München",
    "",
    "Regensburg, 01.10.2026",
    "",
    "Bewerbung als Berechnungsingenieur",
    "",
    "Sehr geehrte Damen und Herren,",
    "Dieser Absatz wird durch den Text aus application.md ersetzt.",
    "Dieser auch.",
    "Mit freundlichen Grüßen",
    "Jonas Michael Becker",
]


def write(root):
    cv_dir = Path(root) / "cv"
    cv_dir.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10.5)
    for text, kind in CV:
        p = d.add_paragraph(style="List Bullet" if kind == "bullet" else None)
        r = p.add_run(text)
        if kind == "name":
            r.bold, r.font.size = True, Pt(20)
        elif kind == "heading":
            r.bold, r.font.size = True, Pt(12)
        elif kind == "bold":
            r.bold = True
    d.save(cv_dir / "Jonas_Becker_Lebenslauf.docx")

    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    for text in LETTER:
        d.add_paragraph(text)
    d.save(cv_dir / "Jonas_Becker_Anschreiben.docx")


if __name__ == "__main__":
    write(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent)
