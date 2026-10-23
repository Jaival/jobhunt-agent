"""Write the sample master CV and cover letter (cv/Lena_Fischer_CV.docx, cv/Lena_Fischer_CoverLetter.docx).

Lena Fischer is fictional: a media and communication student looking for marketing roles. Her CV is a
plain single-column Word document, unlike the engine author's two-column table, to show that the builder
works from cv/layout.yaml and not from one fixed design. The tests in engine/tests build from these files.

Usage: python make_docs.py [workspace folder]   (default: the folder this file is in)
"""
import sys
from pathlib import Path

import docx
from docx.shared import Pt

CV = [
    ("Lena Fischer", "name"),
    ("Marketing Student | Content, Social Media & Analytics", None),
    ("Mannheim | lena.fischer@example.com | +49 151 00000000 | linkedin.com/in/lena-fischer-example", None),
    ("", None),
    ("PROFILE", "heading"),
    ("Media and communication student who plans, writes and measures social media content. Three years of "
     "campus and agency projects, German native, English C1.", None),
    ("", None),
    ("EXPERIENCE", "heading"),
    ("Marketing Intern, Example Media GmbH, Mannheim | 03/2025 – 08/2025", "bold"),
    ("- Planned and published three posts a week on Instagram and LinkedIn for four clients.", None),
    ("- Built a monthly reporting sheet in Google Sheets that the account team still uses.", None),
    ("- Wrote first drafts of two client newsletters a month.", None),
    ("", None),
    ("PROJECTS", "heading"),
    ("Campus Café Campaign", "bold"),
    ("Instagram campaign for the student café: 1,200 new followers in 8 weeks.", None),
    ("Newsletter Relaunch", "bold"),
    ("New layout and sign-up flow for the faculty newsletter; opening rate up from 31% to 44%.", None),
    ("Study Talk Podcast", "bold"),
    ("Six interview episodes with alumni, edited in Audacity and promoted on LinkedIn.", None),
    ("", None),
    ("SKILLS", "heading"),
    ("Content: copywriting, Canva, Adobe Express", None),
    ("Social media: Instagram, LinkedIn, TikTok, Meta Business Suite", None),
    ("Analytics: Google Analytics 4, Google Sheets", None),
    ("Languages: German (native), English (C1)", None),
    ("", None),
    ("EDUCATION", "heading"),
    ("B.A. Media and Communication, Universität Mannheim | 10/2023 – today", "bold"),
]

LETTER = [
    "Lena Fischer",
    "Example Street 1, 68159 Mannheim",
    "lena.fischer@example.com | +49 151 00000000",
    "",
    "Mannheim, 01.10.2026",
    "Example Company GmbH, Mannheim",
    "",
    "Application: Working Student Marketing",
    "",
    "Dear Hiring Team,",
    "This paragraph is replaced by the letter text from application.md.",
    "So is this one.",
    "",
    "Kind regards,",
    "Lena Fischer",
]


def write(root):
    cv_dir = Path(root) / "cv"
    cv_dir.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10.5)
    for text, kind in CV:
        p = d.add_paragraph()
        r = p.add_run(text)
        if kind == "name":
            r.bold, r.font.size = True, Pt(20)
        elif kind == "heading":
            r.bold, r.font.size = True, Pt(12)
        elif kind == "bold":
            r.bold = True
    d.save(cv_dir / "Lena_Fischer_CV.docx")

    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    for text in LETTER:
        d.add_paragraph(text)
    d.save(cv_dir / "Lena_Fischer_CoverLetter.docx")


if __name__ == "__main__":
    write(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent)
