"""Update the engine safely: back up, keep local edits, pull, migrate config.yaml, and prove the documents.

Claude runs these steps for "update the engine" (engine/UPGRADING.md, "Steps for Claude"). Every step prints
what it did and what comes next; none deletes anything without its own command.

Usage (from the workspace folder, the one that contains engine/):
    python engine/tools/upgrade.py status            engine version, new changes on GitHub, local edits, config version
    python engine/tools/upgrade.py prepare           back up config.yaml, CLAUDE.md, custom.md and cv/ to
                                                     backups/pre-upgrade-<date>/, save engine/ as it is (local edits
                                                     included) and list which edits the new engine has a setting for
    python engine/tools/upgrade.py apply             put local edits aside (git stash) and pull the new engine
    python engine/tools/upgrade.py migrate [--dry-run]  bring config.yaml to this engine's version; writes
                                                     cv/layout.yaml from the old CV settings when it's missing
                                                     and adds missing permission rules to .claude/settings.json
                                                     (removes a deny rule for Agent, backup first)
    python engine/tools/upgrade.py compare [--max 5] build a test application (and up to 5 real ones) with the
                                                     saved old engine and the new one, and compare the documents
    python engine/tools/upgrade.py rollback          back to the old engine and config (local edits restored)
    python engine/tools/upgrade.py finish            remove the test builds, the saved engine copy and engine-new/

The first update from an engine without this file runs it from a temporary copy of the new engine
(engine-new/, see UPGRADING.md); it always works on <workspace>/engine.
"""
import argparse
import datetime as dt
import difflib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml

TOOLS = Path(__file__).resolve().parent
ROOT = Path(os.environ.get("JOBHUNT_ROOT") or TOOLS.parents[1])
ENGINE = ROOT / "engine"
sys.path.insert(0, str(TOOLS))
import permissions  # noqa: E402
from setup_check import CONFIG_VERSION  # noqa: E402

# The stock title filter of engines before config_version 2 (search_jobs.py INFRA / SOFTWARE / DROP, with the
# stray character after "tga" fixed), as a profile: block.
STOCK_PROFILE = {
    "role_types": ["Werkstudent", "HiWi", "Internship", "Thesis", "Junior"],
    "include_titles": ["devops", "cloud", "platform", "plattform", "infrastru", r"\bsre\b", "reliabil", "linux",
                       "system ?admin", "systemadmin", "system engineer", "systems engineer", "kubernetes", "network",
                       "netzwerk", "it-admin", r"\bit\b", "informatik", "security", "sicherheit", "software",
                       "backend", "frontend", "full.?stack", "entwickl", "developer", "web", "app", "mobile",
                       "python", "test", r"qa\b", "quality", "automation", "automatisierung", "engineer"],
    "junior_include_titles": ["devops", "cloud", "platform", "plattform", "infrastru", r"\bsre\b", "reliabil",
                              "linux", "system ?admin", "systemadmin", "system engineer", "systems engineer",
                              "kubernetes", "network", "netzwerk", "it-admin", r"\bit\b", "informatik", "security",
                              "sicherheit"],
    "junior_exclude_titles": ["informatik"],
    "exclude_titles": ["senior", r"\blead\b", "principal", "staff", "architect", "architekt", "manager", "leiter",
                       "head of", "director", "vertrieb", "sales", "marketing", "recruit", "buchhalt", "accounting",
                       "finanz", "einkauf", "procurement", "ausbildung", "azubi", "dozent", "logopäd", "lager",
                       "empfang", "bar &", r"service \(", "social media", "design", "elektro", "mechan",
                       "maschinenbau", "gebäude", r"\btga\b", "flight", "unsolicited", "initiativ",
                       "fachinformatiker", "wirtschaftsinformatiker"],
    "score_bonus": [],
}
KNOWN_KEYS = {"config_version", "applicant", "file_prefix", "master_cv", "master_letter", "variants", "letter_city",
              "letter_languages", "name_in_master", "university", "font", "fallback_font", "search", "profile",
              "layout", "internship_duration", "pdf_engine", "scoring", "ledger_expiry_days"}
SEARCH_KEYS = {"home", "areas", "regions", "region_plz", "region_places"}
# Words in a local engine patch -> the setting that does the same job now.
FEATURES = [
    (r"keep_titles|drop_titles", "profile: in config.yaml (migrate writes it from search.keep_titles / drop_titles)"),
    (r"cv_layout", "cv/layout.yaml sections (migrate writes them from cv_layout)"),
    (r"letter_layout", "cv/layout.yaml letter anchors, closing included (migrate writes them from letter_layout)"),
    (r"letter_languages|<!-- de -->|languages:", "letter_languages in config.yaml / languages: [de, en] in application.md"),
    (r"name_in_master", "name_in_master in config.yaml"),
    (r"search\.areas|search_areas|\"areas\"", "search.areas in config.yaml"),
    (r"in_region|regions|region_plz", "search.regions / region_plz / region_places in config.yaml"),
    (r"DATE_LINE|date_line", "cv.date in cv/layout.yaml (the \"City, date\" line on the CV)"),
    (r"internship_months", "internship_duration in config.yaml"),
]
DATE_LINE = re.compile(r"^.+,\s*\d{1,2}\.\s*\d{1,2}\.\s*\d{4}$")
TEST_BODY = (
    "This paragraph is a test of the engine update and is never sent to anyone. It only checks that the letter "
    "keeps its address, date, subject, greeting and closing in the right places after the update. "
) * 7


class UpgradeError(Exception):
    pass


# ---------- helpers ----------

def git(*args, check=True, timeout=120):
    r = subprocess.run(["git", "-C", str(ENGINE), *args], capture_output=True, text=True, encoding="utf-8",
                       timeout=timeout)
    if check and r.returncode:
        raise UpgradeError(f"git {' '.join(args)} failed: {(r.stderr or r.stdout).strip()}")
    return r.stdout.rstrip() if args[:1] == ("status",) else r.stdout.strip()


def edits():
    return [line[3:] for line in git("status", "--porcelain").splitlines() if line.strip()]


TITLE_SCRIPT = """
import json, sys
sys.path.insert(0, sys.argv[1])
import search_jobs
titles = json.load(open(sys.argv[2], encoding="utf-8"))
json.dump({t: bool(search_jobs.in_scope(t)) for t in titles}, open(sys.argv[3], "w", encoding="utf-8"))
"""


def compare_titles(test, ws_old, old_tools):
    """Run the old and the new title filter over every title in data/seen_jobs.csv; report the differences."""
    import csv
    ledger = ROOT / "data" / "seen_jobs.csv"
    if not ledger.exists():
        return ["titles: no data/seen_jobs.csv yet, nothing to compare"], 0
    with ledger.open(encoding="utf-8", newline="") as f:
        titles = sorted({row.get("title", "") for row in csv.DictReader(f)} - {""})
    (test / "titles.json").write_text(json.dumps(titles, ensure_ascii=False), encoding="utf-8")
    kept = {}
    for side, tools, root in (("old", old_tools, ws_old), ("new", ENGINE / "tools", ROOT)):
        out = test / f"titles-{side}.json"
        r = subprocess.run([sys.executable, "-c", TITLE_SCRIPT, str(tools), str(test / "titles.json"), str(out)],
                           cwd=root, env={**os.environ, "JOBHUNT_ROOT": str(root)}, capture_output=True, text=True,
                           encoding="utf-8")
        if r.returncode:
            return [f"titles: the {side} title filter didn't run ({(r.stderr or '').strip().splitlines()[-1:]})"], 0
        kept[side] = json.loads(out.read_text(encoding="utf-8"))
    diff = [t for t in titles if kept["old"][t] != kept["new"][t]]
    lines = [f"titles: {len(titles) - len(diff)} of {len(titles)} ledger titles get the same result"]
    lines += [f"    {'kept before, dropped now' if kept['old'][t] else 'dropped before, kept now'}: {t}" for t in diff]
    return lines, len(diff)


def state_path():
    runs = sorted((ROOT / "backups").glob("pre-upgrade-*/upgrade.json"), key=lambda f: f.stat().st_mtime)
    if not runs:
        raise UpgradeError("no backups/pre-upgrade-*/upgrade.json: run 'upgrade.py prepare' first")
    return runs[-1]


def load_state():
    p = state_path()
    return p, json.loads(p.read_text(encoding="utf-8"))


def save_state(p, state):
    p.write_text(json.dumps(state, indent=1, ensure_ascii=False), encoding="utf-8")


def read_cfg(path=None):
    p = path or ROOT / "config.yaml"
    text = p.read_text(encoding="utf-8")
    return text, yaml.safe_load(text) or {}


def dump(value):
    return yaml.safe_dump(value, allow_unicode=True, sort_keys=False, default_flow_style=None, width=110)


TOP = re.compile(r"^([A-Za-z_][\w-]*)\s*:")


def blocks(text):
    """[(key or None, lines)] for the top-level blocks of a YAML text; comments stay with the block above them."""
    out = [(None, [])]
    for line in text.splitlines(keepends=True):
        m = TOP.match(line)
        if m:
            out.append((m.group(1), [line]))
        else:
            out[-1][1].append(line)
    return out


def set_key(text, key, value, comment="", after=None):
    """Replace a top-level key's block with `key: value`, or add it (after the block `after`, else at the end)."""
    if isinstance(value, (dict, list)) and not (isinstance(value, list) and all(not isinstance(v, (dict, list))
                                                                                 for v in value)):
        new = dump({key: value})
    else:
        flat = yaml.safe_dump(value, allow_unicode=True, default_flow_style=True, width=1000)
        new = f"{key}: {flat.replace('...', '').strip()}\n"
    if comment:
        lines = new.splitlines()
        lines[0] += f"   # {comment}"
        new = "\n".join(lines) + "\n"
    parts = blocks(text)
    for i, (k, lines) in enumerate(parts):
        if k == key:
            trailing = [ln for ln in lines[1:] if ln.strip().startswith("#") or not ln.strip()]
            parts[i] = (k, [new] + trailing)
            break
    else:
        at = next((i + 1 for i, (k, _) in enumerate(parts) if after and k == after), len(parts))
        if at == len(parts) and parts and not "".join(parts[-1][1]).endswith("\n"):
            parts[-1][1].append("\n")
        parts.insert(at, (key, [new]))
    return "".join("".join(lines) for _, lines in parts)


def drop_key(text, key):
    return "".join("".join(lines) for k, lines in blocks(text) if k != key)


# ---------- status / prepare / apply ----------

def cmd_status(a):
    print(f"workspace: {ROOT}")
    if not (ENGINE / ".git").exists():
        raise UpgradeError("engine/ isn't a git clone; clone github.com/Jaival/jobhunt-agent there first")
    v = (ENGINE / "VERSION").read_text(encoding="utf-8-sig").strip() if (ENGINE / "VERSION").exists() else "no VERSION"
    print(f"engine: {v}, {git('log', '-1', '--format=%h %s')} "
          f"(branch {git('rev-parse', '--abbrev-ref', 'HEAD')})")
    fetched = subprocess.run(["git", "-C", str(ENGINE), "fetch", "--quiet"], capture_output=True).returncode == 0
    behind = git("rev-list", "--count", "HEAD..@{u}", check=False) if fetched else ""
    print(f"new changes on GitHub: {behind or 'unknown (offline, or no upstream branch)'}")
    if behind and behind != "0":
        remote = git("show", "@{u}:VERSION", check=False).strip()
        print(f"engine version on GitHub: {remote or 'unknown (no VERSION file there yet)'}")
    if behind and behind != "0":
        print("  " + "\n  ".join(git("log", "--format=%h %s", "HEAD..@{u}").splitlines()[:15]))
    e = edits()
    print(f"local edits in engine/: {len(e)}" + "".join(f"\n  {x}" for x in e))
    _, cfg = read_cfg()
    print(f"config version: {cfg.get('config_version') or 1} (this engine reads {CONFIG_VERSION})")
    wishes = ROOT / "custom.md"
    if wishes.exists():
        print(f"custom.md: {len(re.findall(r'(?m)^- Status: open', wishes.read_text(encoding='utf-8')))} open wish(es)")
    apps = list(ROOT.glob("applications/**/application.md"))
    print(f"applications with application.md: {len(apps)}")
    print(f"workspace is a git repo: {(ROOT / '.git').exists()}")
    return 0


def cmd_prepare(a):
    if not (ENGINE / ".git").exists():
        raise UpgradeError("engine/ isn't a git clone")
    day = f"{dt.date.today():%Y-%m-%d}"
    dest = ROOT / "backups" / f"pre-upgrade-{day}"
    n = 2
    while dest.exists():
        dest = ROOT / "backups" / f"pre-upgrade-{day}-{n}"
        n += 1
    dest.mkdir(parents=True)
    saved = []
    for name in ("config.yaml", "CLAUDE.md", "custom.md"):
        if (ROOT / name).exists():
            shutil.copy2(ROOT / name, dest / name)
            saved.append(name)
    if (ROOT / "cv").is_dir():
        shutil.copytree(ROOT / "cv", dest / "cv")
        saved.append("cv/")
    shutil.copytree(ENGINE, dest / "engine-old", ignore=shutil.ignore_patterns(".git", "__pycache__", ".pytest_cache"))
    changed = edits()
    patch = git("diff", "HEAD")
    (dest / "engine-local-changes.patch").write_text(patch + "\n", encoding="utf-8")
    state = {"date": day, "backup": dest.name, "old_commit": git("rev-parse", "HEAD"),
             "old_branch": git("rev-parse", "--abbrev-ref", "HEAD"), "edits": changed, "stash": None,
             "layout_existed": (ROOT / "cv" / "layout.yaml").exists(), "steps": ["prepare"]}
    save_state(dest / "upgrade.json", state)
    print(f"backup: backups/{dest.name}/ ({', '.join(saved)}, engine-old/ = engine/ as it is now)")
    print(f"old engine: {state['old_commit'][:7]} on {state['old_branch']}")
    if not changed:
        print("local edits in engine/: none")
    else:
        print(f"local edits in engine/: {len(changed)} file(s), saved as backups/{dest.name}/engine-local-changes.patch")
        for f in changed:
            print(f"  {f}")
        found = [setting for rx, setting in FEATURES if re.search(rx, patch)]
        if found:
            print("the new engine has a setting for these edits:")
            for setting in found:
                print(f"  - {setting}")
        print("read the patch for anything else it does: each change without a setting above becomes a wish in "
              "custom.md (engine/templates/workspace/custom.md)")
    print("next: show this to the user; on their yes run 'upgrade.py apply'")
    return 0


def cmd_apply(a):
    p, state = load_state()
    if "apply" in state["steps"]:
        raise UpgradeError("already applied; use 'migrate' next, or 'rollback'")
    if edits():
        git("stash", "push", "--include-untracked", "-m", f"jobhunt upgrade {state['date']}")
        state["stash"] = git("rev-parse", "stash@{0}")
        print("local edits put aside (git stash; 'rollback' brings them back)")
    if state["old_branch"] == "HEAD":
        git("checkout", "main")
    git("pull", "--ff-only", timeout=300)
    state["new_commit"] = git("rev-parse", "HEAD")
    state["steps"].append("apply")
    save_state(p, state)
    print(f"engine: {state['old_commit'][:7]} -> {state['new_commit'][:7]}")
    print("  " + "\n  ".join(git("log", "--format=%h %s", f"{state['old_commit']}..HEAD").splitlines()[:20]))
    print("next: python engine/tools/upgrade.py migrate")
    return 0


# ---------- migrate ----------

def anchor_for(paras, i, what):
    """The shortest anchor that finds paragraph i exactly once."""
    import cv_map
    text = paras[i].text.strip()
    words = text.split()
    tries = [{"starts_with": " ".join(words[:n])} for n in range(1, min(len(words), 6) + 1)] + [{"text": text}]
    for anchor in tries:
        try:
            if cv_map.find(paras, anchor, what)._p is paras[i]._p:
                return anchor
        except cv_map.MapError:
            continue
    raise UpgradeError(f"{what}: '{text}' isn't unique in the master file; write this anchor by hand")


def stock_layout(cfg):
    """The CV map for the stock two-column table CV the engine used to require (cv/README.txt of old engines)."""
    import docx
    import cv_map
    d = docx.Document(str(ROOT / cfg["master_cv"]))
    head = d.tables[0].rows[0].cells[0].paragraphs
    name = next(p.text.strip() for p in head if p.text.strip())
    side = [p.text.strip() for p in d.tables[0].rows[1].cells[-1].paragraphs]
    tech = next((i for i, t in enumerate(side) if cv_map.norm(t) == "technical proficiency"), None)
    skip = next((j - tech - 1 for j in range(tech + 1, len(side)) if ":" in side[j]), 0) if tech is not None else 0
    lay = {"cv": {"headline": {"after": name}, "sections": {
        "projects": {"start": "PROJECTS", "end": ["EXPERIENCE"], "items": "blank_line", "title_marker": "|",
                     "name_split": ","},
        "abilities": {"start": "Skills", "end": ["LEADERSHIP & INITIATIVE"], "items": "no_dash"},
        "technical": {"start": "Technical Proficiency", "skip": skip, "items": "colon", "name_split": ":"}}}}
    return lay, {"recipient": [4, 4], "date": 3, "subject": 5, "greeting": 7, "closing": None}


def flow_layout(cfg):
    """The CV map for a plain-paragraph CV described by an edited engine's config.yaml cv_layout."""
    import docx
    import cv_map
    old = cfg["cv_layout"]
    d = docx.Document(str(ROOT / cfg["master_cv"]))
    paras = d.paragraphs
    heads = list(dict.fromkeys([*(old.get("sections") or []), *[old[k] for k in ("projects", "experience",
                                                                                  "technical") if old.get(k)]]))
    lay = {"cv": {"sections": {}}}
    hi = int(old.get("headline", 1))
    if 0 <= hi < len(paras) and paras[hi].text.strip():
        prev = next((j for j in range(hi - 1, -1, -1) if paras[j].text.strip()), None)
        anchor = {"after": paras[prev].text.strip()} if prev is not None else {"text": paras[hi].text.strip()}
        try:
            ok = cv_map.find(cv_map.paragraphs(d), anchor, "headline")._p is paras[hi]._p
        except cv_map.MapError:
            ok = False
        lay["cv"]["headline"] = anchor if ok else anchor_for(paras, hi, "headline")
    if old.get("date_line", True):
        city = cfg["letter_city"]
        dates = [i for i, p in enumerate(paras) if DATE_LINE.match(p.text.strip()) and p.text.strip().startswith(city)]
        if len(dates) == 1:
            lay["cv"]["date"] = anchor_for(paras, dates[0], "CV date")
    for key, spec in (("projects", {"items": "not_list", "name_split": ["\t", ","]}),
                      ("technical", {"items": "colon", "name_split": ":"})):
        if old.get(key):
            lay["cv"]["sections"][key] = {"start": old[key], "end": [h for h in heads if h != old[key]], **spec}
    return lay, {"recipient": [4, 4], "date": 3, "subject": 5, "greeting": 7, "closing": None,
                 **(cfg.get("letter_layout") or {})}


def letter_map(cfg, pos):
    """Letter anchors from old fixed paragraph positions (0-based)."""
    import docx
    d = docx.Document(str(ROOT / cfg["master_letter"]))
    paras = d.paragraphs
    out = {}
    r0 = pos["recipient"][0] if isinstance(pos.get("recipient"), list) else pos.get("recipient")
    for part, i in (("recipient", r0), ("date", pos.get("date")), ("subject", pos.get("subject")),
                    ("greeting", pos.get("greeting")), ("closing", pos.get("closing"))):
        if i is None or i >= len(paras) or not paras[i].text.strip():
            continue
        if part == "date":
            out[part] = {"starts_with": f"{cfg['letter_city']},"}
        else:
            out[part] = anchor_for(paras, i, f"letter {part}")
    return out


def migrate_1_to_2(text, cfg, notes, dry_run):
    """Config from engines before config_version 2: hardcoded or search.keep_titles title filter, a fixed
    table CV or cv_layout / letter_layout positions, internship_months."""
    search = dict(cfg.get("search") or {})
    if not cfg.get("profile"):
        if search.get("keep_titles"):
            keep = search["keep_titles"]
            drop = search.get("drop_titles")
            profile = {"role_types": STOCK_PROFILE["role_types"],
                       "include_titles": [keep] if isinstance(keep, str) else list(keep),
                       "exclude_titles": ([drop] if isinstance(drop, str) else list(drop)) if drop
                       else STOCK_PROFILE["exclude_titles"],
                       "score_bonus": []}
            notes.append("profile: written from search.keep_titles / drop_titles (same results; every role type "
                         "kept, as before). Narrow role_types if CLAUDE.md says so, after showing the user.")
        else:
            profile = STOCK_PROFILE
            notes.append("profile: the old built-in IT filter, copied exactly. Offer to narrow it to the user's field.")
        text = set_key(text, "profile", profile, "which job titles to keep (engine/tools/scope.py)")
    for k in ("keep_titles", "drop_titles"):
        search.pop(k, None)
    if search != (cfg.get("search") or {}):
        text = set_key(text, "search", search)

    layout_file = ROOT / "cv" / "layout.yaml"
    if not cfg.get("layout"):
        if layout_file.exists():
            notes.append("cv/layout.yaml already exists: kept as it is")
        else:
            try:
                lay, pos = flow_layout(cfg) if cfg.get("cv_layout") else stock_layout(cfg)
                lay["letter"] = letter_map(cfg, pos)
            except UpgradeError:
                raise
            except (KeyError, IndexError, AttributeError, ValueError) as e:  # not the layout the old engine expected
                raise UpgradeError(f"couldn't map the CV automatically ({e}); map it by hand: say 'remap my CV'")
            if not dry_run:
                layout_file.parent.mkdir(exist_ok=True)
                layout_file.write_text("# CV map written by engine/tools/upgrade.py from the old engine's layout "
                                       "(format: engine/tools/cv_map.py)\n" + dump(lay), encoding="utf-8")
            notes.append("cv/layout.yaml: written from " + ("cv_layout / letter_layout" if cfg.get("cv_layout")
                                                            else "the old fixed table layout"))
        text = set_key(text, "layout", "cv/layout.yaml", "CV map (engine/tools/cv_map.py); 'none' = drafts only")
    for k in ("cv_layout", "letter_layout"):
        text = drop_key(text, k)

    if cfg.get("internship_months") and not cfg.get("internship_duration"):
        text = set_key(text, "internship_duration", f"{cfg['internship_months']} months")
    elif not cfg.get("internship_duration"):
        text = set_key(text, "internship_duration", "12 months", "the old engine required this in internship letters")
    text = drop_key(text, "internship_months")
    if not cfg.get("pdf_engine"):
        text = set_key(text, "pdf_engine", "word", "the old engine used Word; auto also falls back to LibreOffice")
    return text


MIGRATIONS = {1: migrate_1_to_2}


def cmd_migrate(a):
    text, cfg = read_cfg()
    version = cfg.get("config_version") or 1
    if version >= CONFIG_VERSION:
        print(f"config.yaml is already version {version}; nothing to migrate")
    notes = []
    while version < CONFIG_VERSION:
        text = MIGRATIONS[version](text, cfg, notes, a.dry_run)
        version += 1
        cfg = yaml.safe_load(text) or {}
    if (cfg.get("config_version") or 1) != version:
        text = drop_key(text, "config_version")
        lines = text.splitlines(keepends=True)
        at = next((i for i, ln in enumerate(lines) if TOP.match(ln)), len(lines))
        lines.insert(at, f"config_version: {version}                               # set by the engine; "
                         "\"update the engine\" migrates older files\n")
        text = "".join(lines)
    cfg = yaml.safe_load(text) or {}
    unknown = sorted(set(cfg) - KNOWN_KEYS) + [f"search.{k}" for k in sorted(set(cfg.get("search") or {}) - SEARCH_KEYS)]
    if a.dry_run:
        print(text)
    else:
        (ROOT / "config.yaml").write_text(text, encoding="utf-8")
        print(f"config.yaml: version {version}")
    for n in notes:
        print(f"  - {n}")
    if unknown:
        print(f"  - not read by this engine (left in place): {', '.join(unknown)}; if they did something, add a "
              "wish to custom.md")
    if not a.dry_run and (ROOT / "cv" / "layout.yaml").exists():
        import cv_map
        problems = cv_map.check(ROOT / "cv" / "layout.yaml", ROOT / cfg["master_cv"], ROOT / cfg["master_letter"],
                                out=lambda s: print(f"  map {s}"))
        print("  CV map: " + ("OK" if not problems else "PROBLEMS: " + "; ".join(problems)))
    try:
        added = permissions.install(ROOT, cfg, dry_run=a.dry_run)
        verb = "would add" if a.dry_run else "added"
        print(f"  permission rules: {verb} {len(added)} to .claude/settings.json" if added
              else "  permission rules: all in place")
        for rule in added:
            print(f"    {rule}")
    except permissions.SettingsError as e:
        print(f"  permission rules: NOT INSTALLED: {e}")
    try:
        p, state = load_state()
        state["steps"].append("migrate")
        save_state(p, state)
    except UpgradeError:
        pass
    print("next: python engine/tools/setup_check.py, then python engine/tools/upgrade.py compare")
    return 0


# ---------- compare ----------

def test_application(cfg, langs, orders):
    meta = {"short": "Upgrade-Test", "role_type": "Junior", "date": f"{dt.date.today():%Y-%m-%d}",
            "recipient": "Test Company GmbH", "subject": "Application: engine update test",
            "greeting": "Dear Test Company team,"}
    if len(langs) > 1:
        meta["languages"] = langs
    cv = {"headline": "Engine update test headline"}
    cv.update(orders)
    if cv:
        meta["cv"] = cv
    body = "\n\n".join([TEST_BODY.strip()] * 1 + ["The second paragraph only checks the order of paragraphs."])
    text = "---\n" + dump(meta) + "---\n"
    if len(langs) > 1:
        text += "".join(f"<!-- {lg} -->\n{body}\n\n" for lg in langs)
    else:
        text += body + "\n"
    return text


def doc_texts(path):
    import docx
    d = docx.Document(str(path))
    out = []
    for el in d.element.body.iter():
        if el.tag.endswith("}p"):
            out.append("".join(t.text or "" for t in el.iter() if t.tag.endswith("}t")))
    return out


def compare_files(old, new):
    za, zb = zipfile.ZipFile(old), zipfile.ZipFile(new)
    parts = sorted((set(za.namelist()) | set(zb.namelist())) - {"docProps/core.xml"})
    differ = [n for n in parts if n not in za.namelist() or n not in zb.namelist() or za.read(n) != zb.read(n)]
    if not differ:
        return "identical", []
    ta, tb = doc_texts(old), doc_texts(new)
    if ta == tb:
        return f"same text, formatting differs ({', '.join(differ[:3])})", []
    diff = [ln for ln in difflib.unified_diff(ta, tb, "old", "new", n=0, lineterm="")
            if ln[:1] in "+-" and not ln.startswith(("+++", "---"))]
    return "text differs", diff[:12]


def build(engine_tools, root, folder):
    env = {**os.environ, "JOBHUNT_ROOT": str(root), "PYTHONIOENCODING": "utf-8"}
    r = subprocess.run([sys.executable, str(engine_tools / "build_application.py"), str(folder), "--no-pdf"],
                       cwd=root, env=env, capture_output=True, text=True, encoding="utf-8")
    return r.returncode, (r.stderr or r.stdout).strip()


def cmd_compare(a):
    import cv_map
    p, state = load_state()
    backup = ROOT / "backups" / state["backup"]
    test = backup / "test"
    shutil.rmtree(test, ignore_errors=True)
    # The old engine runs in a copy of the workspace as it was: the saved config.yaml and cv/, the saved engine.
    ws_old = test / "workspace-old"
    ws_old.mkdir(parents=True)
    for name in ("config.yaml", "CLAUDE.md"):
        if (backup / name).exists():
            shutil.copy2(backup / name, ws_old / name)
    shutil.copytree(backup / "cv", ws_old / "cv")
    shutil.copytree(backup / "engine-old", ws_old / "engine")
    (ws_old / "applications").mkdir()
    old_tools, new_tools = ws_old / "engine" / "tools", ENGINE / "tools"

    _, cfg = read_cfg()
    langs = cfg.get("letter_languages") or ["en"]
    langs = [langs] if isinstance(langs, str) else list(langs)
    orders = {}
    try:
        import docx
        layout = cv_map.load(ROOT / str(cfg.get("layout") or "cv/layout.yaml"))
        master = docx.Document(str(ROOT / cfg["master_cv"]))
        for name, spec in ((layout.get("cv") or {}).get("sections") or {}).items():
            names, _ = cv_map.section(master, name, spec or {})
            if len(names) > 2:   # reorder, but keep the last item last (old engines moved its trailing gap)
                orders[name] = list(reversed(names[:-1])) + names[-1:]
    except Exception as e:
        print(f"note: test application keeps the master order ({e})")

    cases = [("engine update test", test_application(cfg, langs, orders), test_application(cfg, langs, {}), None)]
    real = sorted(ROOT.glob("applications/**/application.md"), key=lambda f: f.parent.name, reverse=True)
    for f in real[:a.max]:
        cases.append((f.parent.name, f.read_text(encoding="utf-8"), None, f.parent / "ad.json"))

    report, worse = [], 0
    for name, text, fallback, ad in cases:
        folder = re.sub(r"[^\w.-]+", "-", name)
        new_dir, old_dir = test / "new" / folder, ws_old / "applications" / folder
        results = {}
        for side, tools, root, d in (("old", old_tools, ws_old, old_dir), ("new", new_tools, ROOT, new_dir)):
            d.mkdir(parents=True)
            (d / "application.md").write_text(text, encoding="utf-8")
            if ad and ad.exists():      # the grounding check accepts the ad's terms
                shutil.copy2(ad, d / "ad.json")
            results[side] = build(tools, root, d)
        if results["old"][0] and fallback:      # the old engine may not know a section: headline only
            for side, tools, root, d in (("old", old_tools, ws_old, old_dir), ("new", new_tools, ROOT, new_dir)):
                for f in d.glob("*.docx"):
                    f.unlink()
                (d / "application.md").write_text(fallback, encoding="utf-8")
                results[side] = build(tools, root, d)
        (ro, eo), (rn, en) = results["old"], results["new"]
        if ro and rn:
            report.append(f"{name}: failed with both engines (old: {eo.splitlines()[-1][:120]} | new: "
                          f"{en.splitlines()[-1][:120]})")
            continue
        if rn:
            worse += 1
            report.append(f"{name}: builds with the old engine but FAILS with the new one: {en[:300]}")
            continue
        if ro:
            report.append(f"{name}: failed with the old engine, builds with the new one ({eo.splitlines()[-1][:120]})")
            continue
        olds = {f.name: f for f in old_dir.glob("*.docx")}
        news = {f.name: f for f in new_dir.glob("*.docx")}
        for fname in sorted(set(olds) | set(news)):
            if fname not in news or fname not in olds:
                worse += 1
                report.append(f"{name}: {fname} only built by the {'old' if fname in olds else 'new'} engine")
                continue
            verdict, diff = compare_files(olds[fname], news[fname])
            if verdict == "text differs":
                worse += 1
            report.append(f"{name}: {fname}: {verdict}" + "".join(f"\n    {ln}" for ln in diff))
    lines, n = compare_titles(test, ws_old, old_tools)
    report += lines
    worse += n
    out = "\n".join(report)
    (backup / "compare-report.txt").write_text(out + "\n", encoding="utf-8")
    print(out)
    print(f"\n{'SAME DOCUMENTS' if not worse else f'{worse} DIFFERENCE(S)'}  (report: backups/{state['backup']}/"
          "compare-report.txt; built files in its test/ folder)")
    state["steps"].append("compare")
    state["compare"] = "same" if not worse else f"{worse} differences"
    save_state(p, state)
    if worse:
        print("next: show the differences to the user. Accept them, fix config.yaml / cv/layout.yaml and compare "
              "again, or run 'upgrade.py rollback'")
    else:
        print("next: setup_check.py must say READY, then 'upgrade.py finish' after the user's yes")
    return 1 if worse else 0


# ---------- rollback / finish ----------

def cmd_rollback(a):
    p, state = load_state()
    backup = ROOT / "backups" / state["backup"]
    if edits():
        git("stash", "push", "--include-untracked", "-m", "jobhunt upgrade: changes made after the update")
        print("changes made in engine/ after the update put aside (git stash)")
    branch = state["old_branch"]
    if branch == "HEAD":
        git("checkout", state["old_commit"])
    else:
        git("checkout", branch)
        git("reset", "--hard", state["old_commit"])
    if state.get("stash"):
        git("stash", "apply", state["stash"])
        print("local edits back in engine/")
    for name in ("config.yaml", "custom.md"):
        if (backup / name).exists():
            shutil.copy2(backup / name, ROOT / name)
    if not state.get("layout_existed") and (ROOT / "cv" / "layout.yaml").exists():
        (ROOT / "cv" / "layout.yaml").rename(backup / "layout.yaml.from-upgrade")
        print(f"cv/layout.yaml moved to backups/{state['backup']}/ (the old engine doesn't use it)")
    state["steps"].append("rollback")
    save_state(p, state)
    print(f"engine back at {state['old_commit'][:7]}; config.yaml restored from backups/{state['backup']}/")
    return 0


def cmd_finish(a):
    p, state = load_state()
    backup = ROOT / "backups" / state["backup"]
    for d in (backup / "test", backup / "engine-old"):
        if d.exists():
            shutil.rmtree(d)
            print(f"removed backups/{state['backup']}/{d.name}/")
    if (ROOT / "engine-new").exists():
        git("worktree", "remove", "--force", str(ROOT / "engine-new"), check=False)
        shutil.rmtree(ROOT / "engine-new", ignore_errors=True)
        git("worktree", "prune", check=False)
        print("removed engine-new/")
    state["steps"].append("finish")
    save_state(p, state)
    print(f"kept: backups/{state['backup']}/ (config.yaml, CLAUDE.md, cv/, the local-edits patch, compare report)")
    if state.get("stash"):
        print(f"your old engine edits are still in git stash {state['stash'][:7]} and in the patch file")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("status", "prepare", "apply", "rollback", "finish"):
        sub.add_parser(name)
    m = sub.add_parser("migrate")
    m.add_argument("--dry-run", action="store_true", help="print the new config.yaml, write nothing")
    c = sub.add_parser("compare")
    c.add_argument("--max", type=int, default=5, help="real applications to compare (default 5)")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    if not shutil.which("git"):
        print("git isn't installed or this session can't run it: open the workspace in Claude Code (the Code tab) "
              "on the computer that has it", file=sys.stderr)
        return 2
    try:
        return globals()[f"cmd_{a.cmd}"](a)
    except UpgradeError as e:
        print(f"STOPPED: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
