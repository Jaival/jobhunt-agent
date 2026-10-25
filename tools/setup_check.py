"""Check whether the workspace is ready for the daily routine, and create the missing scaffolding.

The routine runs this first (engine/ROUTINE.md step 0). When anything is "todo", it switches to the
first-run interview in engine/ONBOARDING.md instead of searching. Unlike the other tools this one
doesn't import config.py, because config.py stops when config.yaml is missing.

Usage (from the workspace folder, the one that contains engine/):
    python engine/tools/setup_check.py            list every check: ok / todo / warn
    python engine/tools/setup_check.py --init     first copy the templates, create the folders and add the
                                                  permission rules to .claude/settings.json (never overwrites
                                                  an existing file; settings are merged)
    python engine/tools/setup_check.py --json     same checks as JSON, for the routine

Exit code 0 = ready (warnings allowed), 1 = something is still "todo".

It prints the engine version first (engine/VERSION and the commit; the routine copies it into the daily brief).
It also looks after the engine itself: a config.yaml from an older engine version is "todo" (say "update the
engine"), local edits in engine/ are a warning (they block updates), and so is a newer engine on GitHub (it runs
`git fetch`, never pull; set JOBHUNT_OFFLINE=1 to skip that).
"""
import argparse
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import permissions
import version
from version import git

ENGINE = Path(__file__).resolve().parents[1]
ROOT = Path(os.environ.get("JOBHUNT_ROOT") or ENGINE.parent)
TEMPLATES = ENGINE / "templates" / "workspace"

FOLDERS = ["applications", "daily", "data", "follow-ups", "interview-prep", "reviews", "backups", "cv", "tracker"]
PACKAGES = {"docx": "python-docx", "openpyxl": "openpyxl", "pypdf": "pypdf", "yaml": "pyyaml", "requests": "requests"}
CONFIG_VERSION = 2   # config.yaml layout this engine reads; engine/tools/upgrade.py migrates older ones
CONFIG_KEYS = ("applicant", "file_prefix", "master_cv", "master_letter", "variants", "letter_city")
# Values that only appear in the templates; a workspace that still has them isn't filled in.
TEMPLATE_VALUES = {"applicant": "Your Name", "file_prefix": "Your_Name",
                   "master_cv": "cv/Your_Name_CV.docx", "master_letter": "cv/Your_Name_CoverLetter.docx"}
PLACEHOLDER = re.compile(r"<(?:Your Name|Name|city|Degree|university|focus|fields|role at company|internship|"
                         r"only ones[^>]*|level|20|40|ASAP|6/12|where[^>]*|tools and skills[^>]*|numbers and[^>]*|"
                         r"languages, tools[^>]*|e\.g\.[^>]*|as soon as possible|\.\.\.|role families[^>]*|"
                         r"preferences[^>]*|your [^>]*|degree|stack|url|year|one-line description|"
                         r"what it does[^>]*|anything that[^>]*|role types)>")
GITIGNORE = "# Public engine repo, cloned here\nengine/\n# Build temp files and search output\n~build_*\n~$*\ndata/candidates.json\n"


def read_config():
    """config.yaml as a dict ({} when it's missing or pyyaml isn't installed yet)."""
    path = ROOT / "config.yaml"
    if not path.exists() or importlib.util.find_spec("yaml") is None:
        return {}
    import yaml
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}


def init():
    """Copy the workspace templates, create the folders and install the permission rules. Returns what it did."""
    done = []
    for src in sorted(TEMPLATES.rglob("*")):
        dst = ROOT / src.relative_to(TEMPLATES)
        if src.is_dir() or src.relative_to(TEMPLATES).parts[0] == ".claude":  # merged by permissions.install
            continue
        if not dst.exists():
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
            done.append(f"copied {dst.relative_to(ROOT).as_posix()}")
    for name in FOLDERS:
        if not (ROOT / name).is_dir():
            (ROOT / name).mkdir(parents=True)
            done.append(f"created {name}/")
    if not (ROOT / ".gitignore").exists():
        (ROOT / ".gitignore").write_text(GITIGNORE, encoding="utf-8")
        done.append("wrote .gitignore (keeps engine/ out of your private repo)")
    try:
        added = permissions.install(ROOT, read_config())
        if added:
            done.append(f"added {len(added)} permission rule(s) to .claude/settings.json")
    except permissions.SettingsError as e:
        done.append(f"permission rules not installed: {e}")
    return done


def placeholders(path):
    text = path.read_text(encoding="utf-8")
    return sorted(set(PLACEHOLDER.findall(text)))


def unmatched_variants(cv_path, variants_path):
    """Ids in variants.yaml whose master text isn't found exactly once in the CV (the builder stops on these)."""
    import docx
    import yaml
    sys.path.insert(0, str(ENGINE / "tools"))
    from cv_map import paragraphs
    key = lambda s: re.sub(r"\s+", " ", s).strip().lower().lstrip("- ")
    paras = [key(p.text) for p in paragraphs(docx.Document(str(cv_path)))]
    variants = yaml.safe_load(variants_path.read_text(encoding="utf-8")) or {}
    return [vid for vid, v in variants.items()
            if sum(1 for p in paras if p and p.startswith(key(v.get("master") or "")[:60])) != 1]


def engine_edits():
    """Files changed in engine/ since its last update (tracked edits and new files)."""
    out = git("status", "--porcelain")
    return [line[3:] for line in (out or "").splitlines() if line.strip()]


def engine_behind():
    """New commits on GitHub that engine/ doesn't have yet (0 when up to date or unknown)."""
    if os.environ.get("JOBHUNT_OFFLINE") or git("fetch", "--quiet") is None:
        return 0
    n = (git("rev-list", "--count", "HEAD..@{u}") or "").strip()
    return int(n) if n.isdigit() else 0


def guard_starts(launcher):
    """Start the hook the way Claude Code does (`python` from PATH) with a Gmail send call: it must exit 2.

    A hook that can't start doesn't block anything (docs/spike-scheduled-tasks.md), so this is the only proof.
    Returns (ok, fix)."""
    python = shutil.which("python")
    if not python:
        return False, ("the guard hook runs `python`, which isn't on PATH (macOS often only has python3): "
                       "install Python from python.org or add a `python` command")
    call = json.dumps({"session_id": "setup-check", "cwd": str(ROOT), "tool_name": "mcp__x__send_message",
                       "tool_input": {}})
    try:
        r = subprocess.run([python, str(launcher)], input=call, capture_output=True, text=True, timeout=30,
                           env={**os.environ, "JOBHUNT_GUARD_SELFTEST": "1", "CLAUDE_PROJECT_DIR": str(ROOT)})
    except (OSError, subprocess.TimeoutExpired) as e:
        return False, f"the guard hook didn't run ({e})"
    if r.returncode == 2 and "guard:" in r.stderr:
        return True, ""
    return False, (f"the guard hook let a Gmail send through (exit {r.returncode}"
                   f"{': ' + r.stderr.strip()[:200] if r.stderr.strip() else ''}): is engine/tools/hooks/guard.py "
                   "there? say 'update the engine'")


def checks():
    """Every check as {area, item, status, fix}. Status: ok, todo (blocks the routine) or warn."""
    out = []

    def add(area, item, ok, fix="", warn=False):
        out.append({"area": area, "item": item, "status": "ok" if ok else ("warn" if warn else "todo"), "fix": fix})

    missing = [pip for mod, pip in PACKAGES.items() if importlib.util.find_spec(mod) is None]
    add("python", "packages", not missing, f"pip install {' '.join(missing)}")

    for name in FOLDERS:
        add("folders", f"{name}/", (ROOT / name).is_dir(), "python engine/tools/setup_check.py --init")

    cfg_path = ROOT / "config.yaml"
    cfg = read_config()
    add("config.yaml", "file", cfg_path.exists(), "python engine/tools/setup_check.py --init")
    if cfg_path.exists():
        version = cfg.get("config_version") or 1
        add("config.yaml", "config version", version >= CONFIG_VERSION,
            f"config.yaml is from an older engine (version {version}, this engine reads {CONFIG_VERSION}): open "
            "this folder in Claude Code and say 'update the engine'")
        for key in CONFIG_KEYS:
            value = cfg.get(key)
            add("config.yaml", key, bool(value) and value != TEMPLATE_VALUES.get(key),
                f"set {key} (onboarding asks for it)")
        home = (cfg.get("search") or {}).get("home") or {}
        add("config.yaml", "search.home", bool(home.get("wo")), "set search.home.wo to your city", warn=True)
        prof = cfg.get("profile")
        inc = (prof or {}).get("include_titles") if isinstance(prof, dict) else None
        add("config.yaml", "profile (which job titles to keep)",
            isinstance(prof, dict) and inc != ["your field"],
            "fill in profile: role_types and include_titles (onboarding round 3)")
        if isinstance(prof, dict):
            sys.path.insert(0, str(ENGINE / "tools"))
            from scope import Scope
            try:
                Scope(prof)
                add("config.yaml", "profile is valid", True)
            except Exception as e:  # bad role type or a broken regex fragment
                add("config.yaml", "profile is valid", False, str(e))

    for key, label in (("master_cv", "master CV"), ("master_letter", "master cover letter")):
        p = ROOT / str(cfg.get(key) or "")
        present = bool(cfg.get(key)) and p.is_file()
        add("cv", label, present, f"put your {label} (.docx) at {cfg.get(key) or 'cv/'}")
        if key == "master_cv" and present and "docx" not in missing and "yaml" not in missing:
            layout = str(cfg.get("layout") or "cv/layout.yaml")
            letter = ROOT / str(cfg.get("master_letter") or "")
            if layout.lower() == "none":
                add("cv", "CV map (cv/layout.yaml)", False,
                    "layout: none: the CV isn't mapped, so tailoring writes drafts only and no documents are built",
                    warn=True)
            elif not (ROOT / layout).is_file():
                add("cv", "CV map (cv/layout.yaml)", False,
                    "map the CV and letter for the builder (onboarding round 5), or set layout: none for drafts only")
            elif letter.is_file():
                sys.path.insert(0, str(ENGINE / "tools"))
                import cv_map
                problems = cv_map.check(ROOT / layout, p, letter, out=lambda _: None)
                add("cv", "CV map matches the master CV and letter", not problems,
                    "; ".join(problems) + " (documents aren't built until this is fixed: say 'remap my CV')",
                    warn=True)
            name = cfg.get("name_in_master")
            if name and name != cfg.get("applicant"):
                import docx
                texts = [r.text for f in (p, letter) if f.is_file()
                         for para in docx.Document(str(f)).paragraphs for r in para.runs]
                add("cv", "name_in_master found in the master files", any(name in t for t in texts),
                    f"'{name}' isn't in one piece in the master CV or letter: fix name_in_master in config.yaml, "
                    "or retype the name in Word", warn=True)
            var_path = ROOT / str(cfg.get("variants") or "")
            if cfg.get("variants") and var_path.is_file():
                left = unmatched_variants(p, var_path)
                add("cv", "variants.yaml matches the master CV", not left,
                    f"not in your CV: {', '.join(left[:6])}; replace these entries with your own CV bullets "
                    "or delete them (onboarding does this)")

    for name, fix in (("CLAUDE.md", "fill in your profile (onboarding asks for it)"),
                      ("projects.md", "describe your projects, or delete the example ones")):
        p = ROOT / name
        if not p.exists():
            add("profile", name, False, "python engine/tools/setup_check.py --init")
            continue
        left = placeholders(p)
        add("profile", name, not left, f"{fix}; still open: {', '.join(left[:6])}")
    claude = ROOT / "CLAUDE.md"
    if claude.exists():
        add("profile", "CLAUDE.md imports engine/RULES.md",
            "@engine/RULES.md" in claude.read_text(encoding="utf-8"), "add the line @engine/RULES.md")

    for name in ("sources.md", "companies.md", "cv/variants.yaml", "tracker/Job_Application_Tracker.xlsx"):
        add("workspace", name, (ROOT / name).exists(), "python engine/tools/setup_check.py --init")

    try:
        left = permissions.missing(ROOT, cfg)
        add("safety", "permission rules (.claude/settings.json)", not left,
            f"{len(left)} rule(s) missing ({', '.join(left[:3])}{', ...' if len(left) > 3 else ''}): "
            "python engine/tools/setup_check.py --init", warn=True)
    except permissions.SettingsError as e:
        add("safety", "permission rules (.claude/settings.json)", False, str(e), warn=True)
    launcher = ROOT / ".claude" / "hooks" / "jobhunt_guard.py"
    if launcher.is_file():
        ok, fix = guard_starts(launcher)
        add("safety", "guard hook starts and blocks", ok, fix, warn=True)

    edits = engine_edits()
    add("engine", "no local edits in engine/", not edits,
        f"engine/ has local changes ({', '.join(edits[:4])}{', ...' if len(edits) > 4 else ''}): they block engine "
        "updates. Say 'update the engine' to keep them safely; write wishes for the engine in custom.md instead",
        warn=True)
    behind = engine_behind()
    add("engine", "engine is up to date", not behind,
        f"an engine update is available ({behind} new change(s)): say 'update the engine'", warn=True)

    import pdf_export
    try:
        engine = pdf_export.pick(cfg.get("pdf_engine"))
        add("documents", f"PDF export: {engine}", engine != "none",
            "pdf_engine: none, so only .docx files are built", warn=True)
        if engine == "libreoffice":
            add("documents", "PDF page counts", False, "LibreOffice's page count can differ from Word's: open "
                "the first CV PDF to check it fits the page limit", warn=True)
    except pdf_export.PdfError as e:
        add("documents", "PDF export", False, f"{e}; the .docx files are still built", warn=True)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--init", action="store_true", help="copy templates and create folders first")
    ap.add_argument("--json", action="store_true", help="print the checks as JSON")
    args = ap.parse_args()

    done = init() if args.init else []
    result = checks()
    todo = [c for c in result if c["status"] == "todo"]
    if args.json:
        print(json.dumps({"root": str(ROOT), "engine_version": version.version(), "engine_commit": version.commit(),
                          "ready": not todo, "init": done, "checks": result}, indent=2))
    else:
        print(f"engine {version.describe()}")
        for line in done:
            print(f"init  {line}")
        for c in result:
            mark = {"ok": "ok  ", "todo": "TODO", "warn": "warn"}[c["status"]]
            fix = f"  -> {c['fix']}" if c["status"] != "ok" and c["fix"] else ""
            print(f"{mark}  {c['area']}: {c['item']}{fix}")
        print(f"\n{'READY' if not todo else f'NOT READY: {len(todo)} item(s) to do'}  ({ROOT})")
    sys.exit(1 if todo else 0)


if __name__ == "__main__":
    main()
