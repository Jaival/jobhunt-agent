"""Add an "Applied" row to the tracker for a sent application folder, then refresh applications.md.

Reads the fields from the folder's job.md (and the contact from outreach.md), backs up the tracker
to backups/<today>/, refuses to run while Excel has the file open, and skips folders already in the
tracker. The tracker only holds sent applications (CLAUDE.md). After saving, the folder moves to
applications/_sent/ (engine/tools/appdirs.py); the folder can be given by name or by any old path.

Usage (from the JobHunt folder):
    python engine/tools/add_applied.py applications/<folder> [--date 2026-10-17] [--via Portal|Email]
                                [--id <job id>] [--note "text"]
"""
import argparse
import copy
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl

import appdirs

from config import ROOT
XLSX = ROOT / "tracker" / "Job_Application_Tracker.xlsx"
LOCK = ROOT / "tracker" / "~$Job_Application_Tracker.xlsx"
GERMAN = ("none", "basic", "B2", "C1+", "unknown")
CADENCE_DAYS = 7


def fields(md):
    out = {}
    for m in re.finditer(r"^\s*-?\s*\*\*([^*:]+):\*\*\s*(.+)$", md, re.M):
        out.setdefault(m.group(1).strip().lower(), m.group(2).strip())
    return out


def first_url(text):
    m = re.search(r"https?://[^\s)>\]]+", text or "")
    return m.group(0) if m else ""


def german(text):
    t = (text or "").lower()
    for g in ("c1+", "b2", "basic", "none"):
        if t.startswith(g.lower()):
            return g if g != "c1+" else "C1+"
    if t.startswith("c1") or t.startswith("c2"):
        return "C1+"
    return "unknown"


def score(text):
    m = re.search(r"\d+", text or "")
    return int(m.group(0)) if m else None


def strip_note(text):
    """'C1+ ("sehr gut"; level disclosed)' -> keep; 'Hamburg (hybrid). About 2 h…' -> first sentence."""
    return re.split(r"(?<=[a-z)])\.\s", text or "", maxsplit=1)[0].strip()


def no_parens(text):
    return re.sub(r"\s*\([^)]*\)", "", text or "").strip()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--date", help="Applied On (YYYY-MM-DD), default today")
    ap.add_argument("--via", default="Portal")
    ap.add_argument("--id", help="Job ID, default: the folder name")
    ap.add_argument("--note", default="")
    a = ap.parse_args()

    folder = appdirs.find(a.folder)
    if not folder:
        print(f"FAILED: no application folder named {a.folder}", file=sys.stderr)
        return 1
    sent_dir = appdirs.BUCKETS["sent"] / folder.name
    job = (folder / "job.md").read_text(encoding="utf-8")
    outreach = (folder / "outreach.md").read_text(encoding="utf-8") if (folder / "outreach.md").exists() else ""
    f = fields(job)
    applied = dt.date.fromisoformat(a.date) if a.date else dt.date.today()
    drafted = dt.date.fromisoformat(folder.name[:10])
    title = re.search(r"^#\s+(.+)$", job, re.M).group(1).strip()
    contact = fields(outreach).get("contact", "")
    if contact.lower().startswith("no contact") or not contact:
        contact = ""
    cv = next(folder.glob("*CV_*.docx"), None)
    letter = next(folder.glob("*Cover_Letter_*.docx"), None)
    # store the paths the folder will have in applications/_sent/
    rel = lambda p: (sent_dir / p.name if p else sent_dir).relative_to(ROOT).as_posix()

    row = {
        "Job ID": a.id or folder.name,
        "Date": dt.datetime.combine(drafted, dt.time()),
        "Status": "Applied",
        "Title": title,
        "Company": no_parens(strip_note(f.get("company", "")).split(",")[0]),
        "Location": strip_note(f.get("location", "")),
        "Source": no_parens(strip_note(f.get("source", ""))),
        "Fit Score": score(f.get("fit score")),
        "German Required": german(f.get("german required")),
        "Match Reasons": f.get("match reasons", ""),
        "Gaps": f.get("gaps", ""),
        "Job URL": first_url(f.get("url")),
        "CV Doc": rel(cv),
        "Cover Letter Doc": rel(letter),
        "Applied Via": a.via,
        "Applied On": dt.datetime.combine(applied, dt.time()),
        "Contact Person": contact,
        "Next Step": f"Follow up if no reply by {applied + dt.timedelta(days=CADENCE_DAYS):%d.%m.%Y}",
        "Notes": a.note,
    }

    if LOCK.exists():
        print("FAILED: the tracker is open in Excel (lock file found). Close it and run again.", file=sys.stderr)
        return 1
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Applications"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    last = max(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value)
    for r in range(2, last + 1):
        if ws.cell(r, col["Job ID"]).value == row["Job ID"] or folder.name in str(ws.cell(r, col["CV Doc"]).value or ""):
            print(f"SKIP: {folder.name} is already in the tracker (row {r})")
            if appdirs.bucket_of(folder) == "drafts":
                appdirs.move(folder, "sent")
                print(f"MOVED {folder.name} -> _sent/")
            return 0
    backup = ROOT / "backups" / f"{dt.date.today():%Y-%m-%d}"
    backup.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, backup / f"Job_Application_Tracker_before-{folder.name[11:40]}.xlsx")
    r = last + 1
    for h, v in row.items():
        c = ws.cell(r, col[h])
        c.value = v
        src = ws.cell(r - 1, col[h])
        if src.has_style:
            c._style = copy.copy(src._style)
    wb.save(XLSX)
    print(f"OK row {r}: {row['Company']} | {row['Title']} | applied {applied:%d.%m.%Y}")
    if appdirs.bucket_of(folder) != "sent":
        appdirs.move(folder, "sent")
        print(f"MOVED {folder.name} -> _sent/")
    for h in ("Company", "Location", "Fit Score", "German Required", "Job URL", "Contact Person"):
        if row[h] in ("", None):
            print(f"  check: '{h}' is empty")
    subprocess.run([sys.executable, str(Path(__file__).resolve().parent / "tracker_to_md.py")], check=True,
                   env={**os.environ, "JOBHUNT_ROOT": str(ROOT)})
    return 0


if __name__ == "__main__":
    sys.exit(main())
