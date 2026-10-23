"""Scan a CV or cover letter for things an ATS (applicant tracking system) parses badly (engine/ats-checklist.md).

This is the mechanical half of the ATS review: it reads the file and reports facts. Claude does the judgment
half (keywords, wording, what to fix) using the checklist, and writes cv/ats-review.md.

Checks on a .docx: text boxes and shapes, text in headers and footers, images, tables and multi-column
sections, fonts, the first line of text, contact details in the body, standard section headings, date formats,
word count. On a .pdf: the extracted text (first line, contact details, headings, dates) and the page count.

Usage (from the workspace folder):
    python engine/tools/ats_scan.py cv/<file>.docx [--letter] [--json]
      --letter  scan as a cover letter (word count of the body, no section headings)
"""
import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HEADINGS = {
    "experience": r"experience|berufserfahrung|work history|employment|praxiserfahrung|erfahrung",
    "education": r"education|ausbildung|studium|bildung|academic",
    "skills": r"skills|kenntnisse|fähigkeiten|kompetenzen|technical proficiency|expertise",
    "languages": r"languages|sprachen|sprachkenntnisse",
}
OPTIONAL_HEADINGS = {
    "projects": r"projects|projekte",
    "certifications": r"certifications|certificates|zertifikate|weiterbildung",
}
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE = re.compile(r"\+?\d[\d ()/-]{7,}\d")
LINKEDIN = re.compile(r"linkedin\.com/in/", re.I)
DATE_FORMS = {
    "MM/YYYY": re.compile(r"\b(0[1-9]|1[0-2])/(19|20)\d\d\b"),
    "MM.YYYY": re.compile(r"\b(0[1-9]|1[0-2])\.(19|20)\d\d\b"),
    "Month YYYY": re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec|januar|februar|märz|mai|juni|"
                             r"juli|oktober|dezember)[a-z]*\.? (19|20)\d\d\b", re.I),
    "YYYY only": re.compile(r"(?<![/.\d])\b(19|20)\d\d\b(?![/.\d])"),
}
COMMON_FONTS = {"arial", "calibri", "cambria", "garamond", "georgia", "helvetica", "open sans", "roboto",
                "lato", "times new roman", "verdana", "aptos", "source sans pro", "noto sans", "segoe ui",
                "google sans", "carlito", "liberation sans", "dejavu sans", "tahoma", "trebuchet ms"}


def _xml(z, name):
    return z.read(name).decode("utf-8", "ignore") if name in z.namelist() else ""


def _text(xml):
    return " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))


def scan_docx(path, letter=False):
    import docx
    findings, facts = [], {}
    with zipfile.ZipFile(path) as z:
        body = _xml(z, "word/document.xml")
        hf = {n: _text(_xml(z, n)).strip() for n in z.namelist() if re.match(r"word/(header|footer)\d*\.xml$", n)}
        fonts = set(re.findall(r'w:ascii="([^"]+)"', body + _xml(z, "word/styles.xml")))
    boxes = len(re.findall(r"<w:txbxContent", body))
    if boxes:
        findings.append(("must-fix", f"{boxes} text box(es) or shape(s) with text: many ATS skip text inside them",
                         "do it yourself: move the text into normal paragraphs or table cells"))
    hf_text = {n: t for n, t in hf.items() if t}
    if hf_text:
        sample = "; ".join(f"{Path(n).stem}: {t[:60]}" for n, t in hf_text.items())
        contact = any(EMAIL.search(t) or PHONE.search(t) for t in hf_text.values())
        findings.append(("must-fix" if contact else "should-fix",
                         f"text in header/footer ({sample}): ATS often ignore headers and footers",
                         "move name and contact details into the first lines of the body"))
    images = len(re.findall(r"<w:drawing|<w:pict", body))
    if images:
        findings.append(("note", f"{images} image(s) or icon(s): ATS ignore them",
                         "make sure every icon has a text label next to it (e.g. 'Phone:'); a photo is optional in Germany"))
    doc = docx.Document(str(path))
    cols = re.findall(r'<w:cols [^>]*w:num="(\d+)"', body)
    if doc.tables or any(int(c) > 1 for c in cols):
        findings.append(("warning", f"layout uses {len(doc.tables)} table(s)"
                         + (f" and {max(int(c) for c in cols)} text columns" if cols else "")
                         + ": most modern ATS read it, some mix up the column order",
                         "your choice: keep the design (record it in CLAUDE.md) or move to one column"))
    odd = sorted(f for f in fonts if f.lower() not in COMMON_FONTS and not f.startswith("+"))
    if odd:
        findings.append(("note", f"uncommon font(s): {', '.join(odd)}",
                         "fine in the PDF; set fallback_font in config.yaml to a common font"))
    facts["fonts"] = sorted(fonts)
    paras = [p.text for p in doc.paragraphs] + [p.text for t in doc.tables for row in t.rows
                                                for c in row.cells for p in c.paragraphs]
    text = "\n".join(paras)
    facts.update(_text_facts(text, letter, findings))
    facts["first_line"] = next((p.strip() for p in paras if p.strip()), "")
    if letter:
        facts["words_total"] = len(re.findall(r"\w+", text))
    return findings, facts


def scan_pdf(path, letter=False):
    import pypdf
    reader = pypdf.PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    findings, facts = [], {"pages": len(reader.pages)}
    if not text.strip():
        findings.append(("must-fix", "no text could be extracted: the PDF is an image (scan or export as picture)",
                         "export from Word or your editor as a text PDF"))
    limit = 1 if letter else 2
    if len(reader.pages) > limit:
        findings.append(("should-fix", f"{len(reader.pages)} pages (aim for {limit} at student/junior level)", "shorten"))
    facts.update(_text_facts(text, letter, findings))
    facts["first_line"] = next((l.strip() for l in text.splitlines() if l.strip()), "")
    return findings, facts


def _text_facts(text, letter, findings):
    facts = {"email": bool(EMAIL.search(text)), "phone": bool(PHONE.search(text)),
             "linkedin": bool(LINKEDIN.search(text))}
    if not facts["email"]:
        findings.append(("must-fix", "no email address in the body text", "add it to the top of the document"))
    if letter:
        return facts
    lines = [l.strip() for l in text.splitlines() if 0 < len(l.strip()) <= 40]
    found = {k: next((l for l in lines if re.fullmatch(rf"(?:{rx})[\s:&a-zäöü-]*", l, re.I)), None)
             for k, rx in {**HEADINGS, **OPTIONAL_HEADINGS}.items()}
    facts["headings"] = {k: v for k, v in found.items() if v}
    missing = [k for k in HEADINGS if not found[k]]
    if missing:
        findings.append(("should-fix", f"no standard heading found for: {', '.join(missing)}",
                         "use a common heading (e.g. 'Experience' / 'Berufserfahrung') so the ATS files the section"))
    forms = Counter({name: len(rx.findall(text)) for name, rx in DATE_FORMS.items()})
    used = [f for f, n in forms.items() if n and f != "YYYY only"]
    facts["date_formats"] = {f: n for f, n in forms.items() if n}
    if len(used) > 1:
        findings.append(("should-fix", f"mixed date formats: {', '.join(used)}", "use one format, e.g. MM/YYYY"))
    return facts


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file")
    ap.add_argument("--letter", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    p = Path(a.file)
    if not p.is_file():
        sys.exit(f"not found: {p}")
    fn = {".docx": scan_docx, ".pdf": scan_pdf}.get(p.suffix.lower())
    if not fn:
        sys.exit("scan a .docx or .pdf (export other formats first)")
    findings, facts = fn(p, a.letter)
    if a.json:
        print(json.dumps({"file": str(p), "findings": [dict(zip(("severity", "issue", "fix"), f)) for f in findings],
                          "facts": facts}, ensure_ascii=False, indent=1))
        return 0
    print(f"ATS scan: {p}")
    for sev, issue, fix in findings or [("ok", "no mechanical problems found", "")]:
        print(f"  [{sev}] {issue}" + (f"\n          -> {fix}" if fix else ""))
    print("Facts: " + json.dumps(facts, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
