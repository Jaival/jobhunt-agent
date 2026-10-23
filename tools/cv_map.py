"""The CV map: where the builder finds the parts of the candidate's own master CV and cover letter.

cv/layout.yaml (path in config.yaml `layout`) names every part the builder may change by text anchors, so the
candidate keeps their own design. Every part is optional: what isn't mapped simply isn't tailored.

    cv:
      headline: {after: "Jane Doe"}          # the line to replace per job: the first non-empty line after the
                                             # line that is exactly "Jane Doe" (or {starts_with: "..."})
      sections:                              # blocks whose items the builder may reorder per job
        projects:                            # cv.projects in application.md; cv.omit_projects works on this one
          start: PROJECTS                    # the heading line, exact text (upper/lower case ignored)
          end: [EXPERIENCE]                  # optional: heading(s) after the section; default: end of its column/cell
          skip: 0                            # optional: lines after the heading that aren't items (default 0)
          items: blank_line                  # where an item starts: blank_line | bold | colon | no_dash | style:<Name>
          title_marker: "|"                  # optional: also starts an item at a line followed by a line containing this
          name_split: ","                    # optional: the item's name is the text before this; default: the whole line
    letter:
      date: {starts_with: "Berlin,"}         # replaced by "<letter_city>, DD.MM.YYYY"
      recipient: {starts_with: "Acme GmbH"}
      subject: {starts_with: "Application"}
      greeting: {starts_with: "Dear"}        # the body runs from the line after it to the first empty line

A section's lines are the paragraphs that follow its heading inside the same container (page body, table cell
or text box), so a section never runs across columns. Experience bullets are found through cv/variants.yaml.

Usage (from the workspace folder):
    python engine/tools/cv_map.py outline <file.docx>   numbered lines with their container, style and bold flag,
                                                        to write layout.yaml from
    python engine/tools/cv_map.py check                 resolve layout.yaml against the master CV and letter and
                                                        print every part and item name (exit 1 on a problem)
"""
import argparse
import re
import sys
from pathlib import Path

import docx
import docx.text.paragraph
import yaml

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
ITEM_RULES = ("blank_line", "bold", "colon", "no_dash")


class MapError(Exception):
    pass


def norm(s):
    return re.sub(r"\s+", " ", s).strip().lower()


def paragraphs(d):
    """Every paragraph in document order (body, tables, text boxes), skipping duplicate fallback copies."""
    out = []
    for p in d.element.body.iter(f"{W}p"):
        if any(a.tag == f"{MC}Fallback" for a in p.iterancestors()):
            continue
        out.append(docx.text.paragraph.Paragraph(p, d._body))
    return out


def container(p):
    """Short label of where a paragraph sits: body, table cell (t1 r2 c1) or text box."""
    el = p._p
    for a in el.iterancestors():
        if a.tag == f"{W}txbxContent":
            return "textbox"
        if a.tag == f"{W}tc":
            tr = a.getparent()
            tbl = tr.getparent()
            body_tbls = [t for t in tbl.getroottree().getroot().iter(f"{W}tbl")]
            return f"t{body_tbls.index(tbl) + 1} r{list(tbl.iterchildren(f'{W}tr')).index(tr) + 1} " \
                   f"c{list(tr.iterchildren(f'{W}tc')).index(a) + 1}"
    return "body"


def is_bold(p):
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold or (r.style is not None and r.style.font.bold) for r in runs)


def find(paras, anchor, what):
    """The one paragraph an anchor names: {after: text}, {starts_with: text} or {text: text}."""
    if isinstance(anchor, str):
        anchor = {"text": anchor}
    if not isinstance(anchor, dict) or len(anchor) != 1:
        raise MapError(f"{what}: anchor must be {{after|starts_with|text: \"...\"}}, got {anchor!r}")
    (kind, value), = anchor.items()
    value = norm(str(value))
    if kind == "after":
        hits = [i for i, p in enumerate(paras) if norm(p.text) == value]
    elif kind == "starts_with":
        hits = [i for i, p in enumerate(paras) if norm(p.text).startswith(value)]
    elif kind == "text":
        hits = [i for i, p in enumerate(paras) if norm(p.text) == value]
    else:
        raise MapError(f"{what}: unknown anchor '{kind}' (use after, starts_with or text)")
    if len(hits) != 1:
        raise MapError(f"{what}: '{value}' found {len(hits)} times (needs exactly 1); the CV may have changed: "
                       "say 'remap my CV'")
    i = hits[0]
    if kind == "after":
        nxt = next((j for j in range(i + 1, len(paras)) if paras[j].text.strip()), None)
        if nxt is None:
            raise MapError(f"{what}: nothing after '{value}'")
        return paras[nxt]
    return paras[i]


def section(d, name, spec):
    """(item names, {name: [p elements]}) for one mapped section."""
    paras = paragraphs(d)
    head = find(paras, {"text": spec.get("start", "")}, f"section {name} start")
    ends = {norm(e) for e in (spec.get("end") or [])}
    sibs = []
    for el in head._p.itersiblings():
        if el.tag != f"{W}p":
            continue
        p = docx.text.paragraph.Paragraph(el, head._parent)
        if norm(p.text) in ends:
            break
        sibs.append(p)
    sibs = sibs[int(spec.get("skip") or 0):]
    rule = spec.get("items", "blank_line")
    marker = spec.get("title_marker")
    if rule not in ITEM_RULES and not str(rule).startswith("style:"):
        raise MapError(f"section {name}: items must be one of {ITEM_RULES} or style:<Name>")

    def starts(i):
        p, t = sibs[i], sibs[i].text.strip()
        if not t:
            return False
        if marker and marker not in t and i + 1 < len(sibs) and marker in sibs[i + 1].text:
            return True
        if rule == "blank_line":
            return i == 0 or not sibs[i - 1].text.strip()
        if rule == "bold":
            return is_bold(p)
        if rule == "colon":
            return ":" in t
        if rule == "no_dash":
            return not t.startswith(("-", "•", "–", "·"))
        return p.style is not None and p.style.name == rule[len("style:"):]

    split = spec.get("name_split")
    blocks, names, current = {}, [], None
    for i, p in enumerate(sibs):
        if starts(i):
            current = p.text.split(split)[0].strip() if split else p.text.strip()
            if current in blocks:
                raise MapError(f"section {name}: two items named '{current}'")
            names.append(current)
            blocks[current] = []
        if current:
            blocks[current].append(p._p)
    blank = lambda e: not "".join(e.itertext()).strip()
    # Blank spacer lines at the end of the section stay put instead of moving with the last item, unless every
    # item ends with a blank line (then it belongs to the item, like the others).
    if names and not (len(names) > 1 and all(blank(blocks[n][-1]) for n in names[:-1])):
        last = blocks[names[-1]]
        while len(last) > 1 and blank(last[-1]):
            last.pop()
    if not names:
        raise MapError(f"section {name}: no items found after '{spec.get('start')}'")
    return names, blocks


def load(path):
    if not Path(path).is_file():
        raise MapError(f"{path} not found: map the CV first (onboarding round 5, or say 'remap my CV')")
    m = yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}
    for key in m:
        if key not in ("cv", "letter"):
            raise MapError(f"layout.yaml: unknown top-level key '{key}' (use cv and letter)")
    return m


def letter_parts(d, lmap):
    """{part: paragraph} for date, recipient, subject, greeting, plus 'body' = list of body paragraphs."""
    paras = paragraphs(d)
    parts = {k: find(paras, lmap[k], f"letter {k}") for k in ("date", "recipient", "subject", "greeting")
             if lmap.get(k)}
    if "greeting" not in parts:
        raise MapError("letter: greeting anchor is required (the body starts after it)")
    body = []
    for el in parts["greeting"]._p.itersiblings():
        if el.tag != f"{W}p":
            continue
        p = docx.text.paragraph.Paragraph(el, parts["greeting"]._parent)
        if not p.text.strip():
            if body:
                break
            continue
        body.append(p)
    if not body:
        raise MapError("letter: no body paragraphs after the greeting")
    parts["body"] = body
    return parts


def check(layout_path, cv_path, letter_path, out=print):
    """Resolve everything; returns a list of problems (empty = fine)."""
    problems = []
    try:
        m = load(layout_path)
    except MapError as e:
        return [str(e)]
    cvm = m.get("cv") or {}
    d = docx.Document(str(cv_path))
    if cvm.get("headline"):
        try:
            out(f"headline: {find(paragraphs(d), cvm['headline'], 'headline').text.strip()}")
        except MapError as e:
            problems.append(str(e))
    for name, spec in (cvm.get("sections") or {}).items():
        try:
            names, _ = section(d, name, spec or {})
            out(f"section {name}: {names}")
        except MapError as e:
            problems.append(str(e))
    if m.get("letter"):
        try:
            parts = letter_parts(docx.Document(str(letter_path)), m["letter"])
            for k, p in parts.items():
                out(f"letter {k}: " + (f"{len(p)} paragraphs" if k == "body" else p.text.strip()[:70]))
        except MapError as e:
            problems.append(str(e))
    return problems


def outline(path):
    d = docx.Document(str(path))
    for i, p in enumerate(paragraphs(d)):
        style = p.style.name if p.style is not None else ""
        flags = "B" if is_bold(p) else " "
        print(f"{i:4} {container(p):10} {flags} {style[:14]:14} {p.text.strip()[:90]}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    o = sub.add_parser("outline")
    o.add_argument("file")
    sub.add_parser("check")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    if a.cmd == "outline":
        outline(a.file)
        return 0
    from config import CFG, path
    if str(CFG.get("layout", "cv/layout.yaml")).lower() == "none":
        print("layout: none (CV not mapped; tailoring writes drafts only)")
        return 0
    problems = check(path("layout") if CFG.get("layout") else path("master_cv").parent / "layout.yaml",
                     path("master_cv"), path("master_letter"))
    for p in problems:
        print(f"PROBLEM {p}")
    print("OK" if not problems else f"{len(problems)} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
