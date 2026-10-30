"""Build <prefix>_CV_<short>.docx, <prefix>_Cover_Letter_<short>.docx and their PDFs from an application.md.

application.md = YAML frontmatter (CV choices, letter header) + the letter body in Markdown
(one paragraph per block, blank line between paragraphs). See engine/templates/application_template.md.

The master CV and master cover letter stay the templates (CLAUDE.md rule 5): the builder only
changes the headline (and a mapped "City, date" line), reorders the items of the sections named in the CV map
(cv/layout.yaml, see cv_map.py), and swaps bullet text for approved variants from cv/variants.yaml. In the letter
it fills the date, the address block (recipient: one line, or a list / several lines), subject, greeting and body.
It fails before writing anything if a check breaks.

Letters in German and/or English: `languages: [de, en]` in the frontmatter (default: config.yaml
`letter_languages`, else [en]). With two languages the body has a `<!-- de -->` and an `<!-- en -->` part, and
recipient / subject / greeting / closing may be {de: ..., en: ...}; the mapped closing line gets each language's
closing. One letter per language: <prefix>_Cover_Letter_<short>_DE.docx / _EN.docx (one language keeps the plain
name). config.yaml `name_in_master`: the name as the master files spell it, replaced by `applicant` in every
document (for example a middle name you leave out).
Every CV keeps all projects; cv.omit_projects only works together with cv.omit_approved: <date>
(the applicant's approval). After the PDF export (pdf_export.py, config pdf_engine) it fails if the CV is
over 2 pages or the letter over 1.

Grounding: every number and every skill or tool term (tools/skills.txt plus config.yaml profile.score_bonus) in a
letter or the CV headline must appear in the candidate's own files (master CV, CLAUDE.md without its Gaps section,
projects.md, cv/variants.yaml) or in the job's ad.json; otherwise the build fails with the list. Terms only the ad
has pass with a warning. `grounding_ok: [term, ...]` in the frontmatter lets reviewed exceptions through. Numbers
written as words ("three") and the context of a term (claimed or named as a gap) aren't checked.

Output for ATS autofill: file names start with the applicant's name (parsers read names from file
names), document metadata carries the applicant as author and a neutral title (Word copies it into
the PDF), and every font is set to FONT with FALLBACK_FONT as the substitute named in the font table.

Usage (from the JobHunt folder):
    python engine/tools/build_application.py applications/<folder> [--no-pdf]   (a bare folder name also works)
"""
import argparse
import copy
import datetime as dt
import re
import sys
from pathlib import Path

import docx
import docx.text.paragraph
import yaml

import appdirs
import cv_map
import pdf_export
from cv_map import MapError, paragraphs

from config import CFG, ROOT, path

MASTER_CV = path("master_cv")
MASTER_LETTER = path("master_letter")
VARIANTS = path("variants")
LAYOUT = str(CFG.get("layout") or "cv/layout.yaml")
CV_KEYS = {"headline", "bullets", "variants", "omit_projects", "omit_approved"}  # plus the mapped section names

STUDENT_TYPES = {"Werkstudent", "HiWi", "Internship"}
WORDS_MIN, WORDS_MAX = 250, 350
CV_MAX_PAGES, LETTER_MAX_PAGES = 2, 1
BANNED = [
    "passionate", "i am writing to express", "just following up", "just checking in",
    "touching base", "circling back", "team player", "synergy", "hit the ground running",
    "dynamic environment", "results-driven", "go-getter",
    # German letters
    "hiermit bewerbe ich mich", "mit großem interesse", "teamplayer", "hochmotiviert",
]
LANGS = ("de", "en")
CLOSING = {"de": "Mit freundlichen Grüßen", "en": "Kind regards"}
# Phrases the student-letter checks accept per language (RULES.md rule 9).
START_WORDS = {"en": ["as soon as possible", "immediately"],
               "de": ["ab sofort", "so bald wie möglich", "sobald wie möglich", "schnellstmöglich", "zum nächstmöglichen"]}
HOURS_WORDS = {"en": ["hours per week", "hours a week"],
               "de": ["stunden pro woche", "wochenstunden", "stunden wöchentlich", "stunden in der woche"]}
APPLICANT = CFG["applicant"]          # name in metadata; files start with FILE_PREFIX
FILE_PREFIX = CFG["file_prefix"]
UNIVERSITY = CFG.get("university", "")   # student letters must name it
INTERNSHIP_DURATION = CFG.get("internship_duration") or ""   # internship letters must state it: "12 months" or
                                                              # {de: "12 Monate", en: "12 months"}
NAME_IN_MASTER = CFG.get("name_in_master") or ""   # the name as the master files spell it, replaced by APPLICANT
FONT = CFG.get("font", "Open Sans")                # font for CV and letter
FALLBACK_FONT = CFG.get("fallback_font", "Open Sans")  # substitute when FONT isn't installed on the reader's machine
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class BuildError(Exception):
    pass


# ---------- application.md ----------

def read_application(folder):
    path = folder / "application.md"
    if not path.exists():
        raise BuildError(f"{path} not found")
    text = path.read_text(encoding="utf-8")
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", text, re.S)
    if not m:
        raise BuildError("application.md needs a YAML frontmatter block between --- lines")
    meta = yaml.safe_load(m.group(1)) or {}
    for key in ("short", "role_type", "recipient", "subject", "greeting"):
        if not meta.get(key):
            raise BuildError(f"frontmatter is missing '{key}'")
    langs = meta.get("languages") or CFG.get("letter_languages") or ["en"]
    langs = [langs] if isinstance(langs, str) else list(langs)
    bad = [lg for lg in langs if lg not in LANGS]
    if bad:
        raise BuildError(f"languages: {bad} unknown; use {list(LANGS)}")
    # The body splits at <!-- de --> / <!-- en --> lines; with one language, text without a marker is its body.
    pieces = re.split(r"^\s*<!--\s*(de|en)\s*-->\s*$", m.group(2), flags=re.M | re.I)
    raw = {None: pieces[0]}
    for i in range(1, len(pieces), 2):
        raw[pieces[i].lower()] = pieces[i + 1]
    bodies = {}
    for lg in langs:
        text = raw.get(lg, raw[None] if len(langs) == 1 else None)
        if text is None:
            raise BuildError(f"no letter body for '{lg}': put a '<!-- {lg} -->' line before it")
        paras = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", text)]
        bodies[lg] = [p for p in paras if p and not p.startswith("<!--")]
        if not bodies[lg]:
            raise BuildError(f"the letter body for '{lg}' is empty")
        for key in ("recipient", "subject", "greeting"):
            pick_lang(meta[key], lg, key)
    return meta, bodies


def pick_lang(value, lang, what=""):
    """A frontmatter value that is the same for every language, or {de: ..., en: ...}."""
    if isinstance(value, dict):
        if not value.get(lang):
            raise BuildError(f"'{what}' has no '{lang}' value")
        return value[lang]
    return value


# ---------- docx helpers ----------

def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting (or the '-' bullet run)."""
    runs = p.runs
    # A "- " bullet whose dash shares a run with the text keeps its dash; without it the
    # no_dash item rule (cv_map.py) would treat the bullet as an item start and move it.
    prefix = "- " if p.text.lstrip().startswith("-") and not text.lstrip().startswith("-") else ""
    if not runs:
        p.add_run(prefix + text)
        return
    if runs[0].text.strip() == "-" and len(runs) > 1:
        runs[1].text = " " + text
        rest = runs[2:]
    else:
        runs[0].text = prefix + text
        rest = runs[1:]
    for r in rest:
        r.text = ""


def reorder_blocks(blocks, order, names, what):
    """blocks: {name: [p elements]} in document order, contiguous. order: list of names."""
    if order is None:
        return []
    unknown = [n for n in order if n not in blocks]
    if unknown:
        raise BuildError(f"unknown {what}: {unknown}; known: {list(blocks)}")
    missing = [n for n in names if n not in order]
    final = list(order) + missing  # anything not named keeps its relative place at the end
    last = blocks[names[-1]][-1]
    anchor, parent = last.getnext(), last.getparent()
    for n in names:
        for e in blocks[n]:
            parent.remove(e)
    for n in final:
        for e in blocks[n]:
            if anchor is not None:
                anchor.addprevious(e)
            else:
                parent.append(e)
    return final


def load_layout():
    if LAYOUT.lower() == "none":
        raise BuildError("config.yaml has layout: none, so the CV isn't mapped and documents can't be built; "
                         "send the drafts by hand or say 'remap my CV'")
    try:
        return cv_map.load(ROOT / LAYOUT)
    except MapError as e:
        raise BuildError(str(e))


# ---------- CV ----------

def build_cv(meta, variants, layout, out):
    cvm = meta.get("cv") or {}
    lmap = layout.get("cv") or {}
    sections = lmap.get("sections") or {}
    unknown = [k for k in cvm if k not in CV_KEYS and k not in sections]
    if unknown:
        raise BuildError(f"unknown cv keys {unknown}; mapped sections: {list(sections)} (cv/layout.yaml)")
    d = docx.Document(MASTER_CV)
    summary = []

    if cvm.get("headline"):
        if not lmap.get("headline"):
            raise BuildError("cv.headline is set but cv/layout.yaml maps no headline")
        try:
            set_text(cv_map.find(paragraphs(d), lmap["headline"], "headline"), cvm["headline"])
        except MapError as e:
            raise BuildError(str(e))
        summary.append(f"headline: {cvm['headline']}")

    if lmap.get("date"):
        try:
            set_text(cv_map.find(paragraphs(d), lmap["date"], "CV date"), f"{CFG['letter_city']}, {letter_date(meta):%d.%m.%Y}")
        except MapError as e:
            raise BuildError(str(e))

    # Variants and bullet order use ids from cv/variants.yaml, located by their master text.
    by_id = {}
    paras = paragraphs(d)
    for vid, v in variants.items():
        hits = [p for p in paras if key(p.text) and key(p.text).startswith(key(v["master"])[:60])]
        if len(hits) != 1:
            raise BuildError(f"variant '{vid}': master text found {len(hits)} times in the master CV")
        by_id[vid] = hits[0]
    picks = {vid: v["default"] for vid, v in variants.items() if v.get("default")}
    for vid in cvm.get("variants") or {}:
        if vid not in variants:
            raise BuildError(f"unknown variant id '{vid}'")
    picks.update(cvm.get("variants") or {})
    for vid, pick in picks.items():
        if pick == "master":
            continue
        alt = (variants[vid].get("variants") or {}).get(pick)
        if not alt:
            raise BuildError(f"variant '{vid}' has no option '{pick}'")
        if not alt.get("approved"):
            raise BuildError(f"variant '{vid}.{pick}' is not approved yet (cv/variants.yaml)")
        set_text(by_id[vid], alt["text"])
        summary.append(f"variant {vid}: {pick}")

    for group, order in (cvm.get("bullets") or {}).items():
        ids = [vid for vid, v in variants.items() if v.get("group") == group]
        if not ids:
            raise BuildError(f"unknown bullet group '{group}'")
        blocks = {vid: [by_id[vid]._p] for vid in ids}
        ids.sort(key=lambda vid: list(by_id[vid]._p.getparent()).index(by_id[vid]._p))
        reorder_blocks(blocks, [f"{group}.{n}" for n in order], ids, f"bullet in '{group}'")
        summary.append(f"{group} bullets: {', '.join(order)}")

    # Every CV keeps all projects (CLAUDE.md rule 5); omit_projects only works with the applicant's dated approval.
    omit = cvm.get("omit_projects") or []
    if omit and "projects" not in sections:
        raise BuildError("omit_projects needs a section named 'projects' in cv/layout.yaml")
    for name, spec in sections.items():
        try:
            names, blocks = cv_map.section(d, name, spec or {})
        except MapError as e:
            raise BuildError(str(e))
        if name == "projects" and omit:
            approved = cvm.get("omit_approved") or meta.get("omit_approved")
            if not approved or not re.match(r"^\d{4}-\d{2}-\d{2}$", str(approved)):
                raise BuildError("omit_projects needs omit_approved: <YYYY-MM-DD>, set only after the applicant approved "
                                 "leaving these projects out of this CV")
            unknown = [n for n in omit if n not in blocks]
            if unknown:
                raise BuildError(f"omit_projects: unknown project: {unknown}; known: {names}")
            both = [n for n in omit if n in (cvm.get("projects") or [])]
            if both:
                raise BuildError(f"project listed in both projects and omit_projects: {both}")
            for n in omit:
                for e in blocks.pop(n):
                    e.getparent().remove(e)
                names.remove(n)
            summary.append(f"omitted projects (approved {approved}): {', '.join(omit)}")
        if cvm.get(name):
            reorder_blocks(blocks, cvm[name], names, f"{name} item")
            summary.append(f"{name}: {', '.join(cvm[name])}")

    d.save(out)
    return d, summary


def letter_date(meta):
    date = meta.get("date") or dt.date.today()
    return dt.date.fromisoformat(date) if isinstance(date, str) else date


def norm(s):
    return re.sub(r"\s+", " ", s).strip().lower()


def key(s):
    """Match text without the '- ' bullet prefix."""
    return norm(s).lstrip("- ")


# ---------- letter ----------

def lines_of(value):
    """Address lines from a frontmatter value: a list, or a string with one line per line."""
    lines = [str(x) for x in value] if isinstance(value, list) else str(value).splitlines()
    return [x.strip() for x in lines if x.strip()]


def fill_recipient(block, value):
    """Write the address into the master's address block: one paragraph per line, or line breaks inside a
    one-paragraph block. Extra master lines are removed, missing ones added in the same format."""
    lines = lines_of(value)
    if len(block) == 1:
        set_text(block[0], "\n".join(lines))
        return
    for p, text in zip(block, lines):
        set_text(p, text)
    anchor = block[-1]._p
    for text in lines[len(block):]:
        new = copy.deepcopy(anchor)
        anchor.addnext(new)
        anchor = new
        set_text(docx.text.paragraph.Paragraph(new, block[-1]._parent), text)
    for p in block[len(lines):]:
        p._p.getparent().remove(p._p)


def build_letter(meta, body, lang, langs, layout, out):
    d = docx.Document(MASTER_LETTER)
    if not layout.get("letter"):
        raise BuildError("cv/layout.yaml maps no letter (date, recipient, subject, greeting)")
    try:
        parts = cv_map.letter_parts(d, layout["letter"])
    except MapError as e:
        raise BuildError(str(e))
    for part, text in (("date", f"{CFG['letter_city']}, {letter_date(meta):%d.%m.%Y}"),
                       ("subject", pick_lang(meta["subject"], lang, "subject")),
                       ("greeting", pick_lang(meta["greeting"], lang, "greeting"))):
        if part in parts:
            set_text(parts[part], text)
    if "recipient" in parts:
        fill_recipient(parts["recipient"], pick_lang(meta["recipient"], lang, "recipient"))
    # The closing line keeps the master text unless application.md sets one or there are several languages.
    if "closing" in parts and (meta.get("closing") or len(langs) > 1):
        set_text(parts["closing"], pick_lang(meta.get("closing") or CLOSING, lang, "closing"))
    elif meta.get("closing"):
        raise BuildError("closing is set but cv/layout.yaml maps no letter closing")
    # New body paragraphs copy the first master body paragraph; an empty line between master paragraphs is kept
    # between the new ones too.
    old = [p._p for p in parts["body_all"]]
    proto, anchor = parts["body"][0]._p, old[-1]
    spacer = parts["spacer"]._p if parts["spacer"] is not None else None
    parent = parts["body"][0]._parent
    for i, text in enumerate(body):
        if i and spacer is not None:
            gap = copy.deepcopy(spacer)
            anchor.addnext(gap)
            anchor = gap
        new = copy.deepcopy(proto)
        anchor.addnext(new)
        anchor = new
        set_text(docx.text.paragraph.Paragraph(new, parent), text)
    for e in old:
        e.getparent().remove(e)
    d.save(out)
    return d


# ---------- fonts and metadata ----------

def apply_fonts(d):
    """Set every explicit font (runs, styles, defaults, numbering) to FONT and name FALLBACK_FONT as its substitute."""
    parts = [d.element, d.styles.element]
    numbering = getattr(d.part, "numbering_part", None)
    try:
        if numbering is not None:
            parts.append(numbering.element)
    except Exception:
        pass
    for root in parts:
        for rf in root.iter(f"{W}rFonts"):
            for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
                rf.attrib.pop(f"{W}{attr}", None)
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                rf.set(f"{W}{attr}", FONT)
    defaults = d.styles.element.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
    if defaults is not None and defaults.find(f"{W}rFonts") is None:
        rf = defaults.makeelement(f"{W}rFonts", {f"{W}{a}": FONT for a in ("ascii", "hAnsi", "cs", "eastAsia")})
        defaults.insert(0, rf)
    font_table = next((rel.target_part for rel in d.part.rels.values() if rel.reltype.endswith("/fontTable")), None)
    if font_table is not None:
        from lxml import etree
        root = etree.fromstring(font_table.blob)
        entry = next((f for f in root.iter(f"{W}font") if f.get(f"{W}name") == FONT), None)
        if entry is None:  # a complete entry; a bare one made Word hang when opening the letter
            entry = etree.SubElement(root, f"{W}font", {f"{W}name": FONT})
            for tag, val in (("charset", "00"), ("family", "swiss"), ("pitch", "variable")):
                etree.SubElement(entry, f"{W}{tag}", {f"{W}val": val})
        for old in entry.findall(f"{W}altName"):
            entry.remove(old)
        alt = etree.Element(f"{W}altName", {f"{W}val": FALLBACK_FONT})
        entry.insert(0, alt)
        font_table._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def documents(d):
    """Every document part with paragraphs: the body and each section's headers and footers."""
    yield d
    for sec in d.sections:
        for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
                   sec.even_page_header, sec.even_page_footer):
            if not hf.is_linked_to_previous:
                yield hf


def rename(d):
    """Replace NAME_IN_MASTER with APPLICANT run by run, so formatting stays. Returns the number of changes."""
    n = 0
    if not NAME_IN_MASTER or NAME_IN_MASTER == APPLICANT:
        return n
    for part in documents(d):
        root = part.element if part is d else part._element
        for el in root.iter(f"{W}p"):
            for r in docx.text.paragraph.Paragraph(el, part).runs:
                if NAME_IN_MASTER in r.text:
                    r.text = r.text.replace(NAME_IN_MASTER, APPLICANT)
                    n += 1
            if NAME_IN_MASTER not in APPLICANT and NAME_IN_MASTER in "".join(t.text or "" for t in el.iter(f"{W}t")):
                raise BuildError(f"name_in_master '{NAME_IN_MASTER}' is split over several formatting runs in the "
                                 "master; retype it in one go in Word")
    return n


def set_metadata(d, kind):
    cp = d.core_properties
    cp.author = APPLICANT
    cp.last_modified_by = APPLICANT
    cp.title = f"{APPLICANT} - {kind}"
    cp.subject = cp.keywords = cp.comments = cp.category = ""


# ---------- checks ----------

def doc_text(d):
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts += [p.text for p in c.paragraphs]
    return "\n".join(parts)


def internship_duration(lang):
    """The duration an internship letter in this language must state ("" = no check)."""
    if isinstance(INTERNSHIP_DURATION, dict):
        return str(INTERNSHIP_DURATION.get(lang) or "")
    text = str(INTERNSHIP_DURATION)
    m = re.fullmatch(r"\s*(\d+)\s*months?\s*", text, re.I)
    return f"{m.group(1)} Monate" if lang == "de" and m else text


def check(meta, bodies, cv_doc, letter_docs):
    problems = []
    tag = lambda lg: "letter" if len(bodies) == 1 else f"letter {lg.upper()}"
    for name, d in [("CV", cv_doc)] + [(tag(lg), d) for lg, d in letter_docs.items()]:
        txt = doc_text(d)
        if "[" in txt or "]" in txt:
            problems.append(f"{name}: contains '[' or ']' (placeholder?)")
        low = txt.lower()
        problems += [f"{name}: banned phrase '{b}'" for b in BANNED if b in low]
    counts = {}
    for lg, body in bodies.items():
        words = counts[lg] = sum(len(p.split()) for p in body)
        if not WORDS_MIN <= words <= WORDS_MAX:
            problems.append(f"{tag(lg)} body has {words} words (allowed {WORDS_MIN}-{WORDS_MAX})")
        if meta["role_type"] in STUDENT_TYPES:
            letter = " ".join(body).lower()
            if UNIVERSITY and UNIVERSITY.lower() not in letter:
                problems.append(f"{tag(lg)}: student letter must mention {UNIVERSITY} enrolment")
            if not any(w in letter for w in START_WORDS[lg]):
                problems.append(f"{tag(lg)}: student letter must state the start ('{START_WORDS[lg][0]}')")
            if not any(w in letter for w in HOURS_WORDS[lg]):
                problems.append(f"{tag(lg)}: student letter must state the {HOURS_WORDS[lg][0]}")
            months = internship_duration(lg)
            if meta["role_type"] == "Internship" and months and months.lower() not in letter:
                problems.append(f"{tag(lg)}: internship letter must mention the duration ({months})")
    return problems, counts


# ---------- grounding ----------

NUMBER = re.compile(r"(?<![\w.,])(\d+(?:[.,]\d+)*)(?!\w)(\s*(?:%|prozent\b|percent\b|per cent\b))?", re.I)


def numbers(text):
    """The numbers in a text, normalised: "1,200" / "1.200" -> 1200, "2,5" -> 2.5, "18 Prozent" -> 18%; other
    forms (dates) as written."""
    found = set()
    for n, percent in NUMBER.findall(text):
        if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", n):
            n = re.sub(r"[.,]", "", n)
        elif re.fullmatch(r"\d+[.,]\d+", n):
            n = n.replace(",", ".")
        found.add((str(int(n)) if n.isdigit() else n) + ("%" if percent else ""))
    return found


def grounded_number(n, sources):
    """A percentage needs the same percentage in the sources; a plain number matches either form."""
    return n in sources or (not n.endswith("%") and f"{n}%" in sources)


def load_terms():
    """Skill and tool terms from tools/skills.txt and config.yaml profile.score_bonus: (term, pattern) pairs."""
    lines = (Path(__file__).with_name("skills.txt").read_text(encoding="utf-8").splitlines()
             + [str(t) for t in (CFG.get("profile") or {}).get("score_bonus") or []])
    terms = {}
    for line in lines:
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        exact = line.startswith("=")
        term = line.lstrip("=")
        terms[term.lower()] = (term, re.compile(rf"(?<!\w){re.escape(term)}(?!\w)", 0 if exact else re.I))
    return list(terms.values())


def own_sources(variants):
    """The candidate's own files as one text: the master CV, CLAUDE.md without its Gaps section, projects.md, the
    master and approved texts of cv/variants.yaml, and the internship duration from config.yaml."""
    parts = [doc_text(docx.Document(MASTER_CV)), str(INTERNSHIP_DURATION)]
    claude = ROOT / "CLAUDE.md"
    if claude.exists():
        sections = re.split(r"^(?=## )", claude.read_text(encoding="utf-8"), flags=re.M)
        parts += [s for s in sections if not re.match(r"## .*(gap|never claim|lücke)", s, re.I)]
    projects = ROOT / "projects.md"
    if projects.exists():
        parts.append(projects.read_text(encoding="utf-8"))
    for v in variants.values():
        parts.append(v.get("master", ""))
        parts += [alt.get("text", "") for alt in (v.get("variants") or {}).values() if alt.get("approved")]
    return "\n".join(parts)


def grounding(meta, bodies, variants, folder):
    """Every number and every skill term (load_terms) in the letters and the CV headline must appear in the
    candidate's own files or the job's ad.json. Returns (problems, warnings); warnings name terms only the ad has."""
    own = own_sources(variants)
    ad_file = folder / "ad.json"
    ad = ad_file.read_text(encoding="utf-8") if ad_file.exists() else ""
    own_numbers, ad_numbers = numbers(own), numbers(ad)
    terms = load_terms()
    ok = {str(t).lower() for t in meta.get("grounding_ok") or []}
    ok |= {n for t in ok for n in numbers(t)}
    texts = {f"letter {lg.upper()}" if len(bodies) > 1 else "letter": "\n".join(b) for lg, b in bodies.items()}
    headline = (meta.get("cv") or {}).get("headline")
    if headline:
        texts["CV headline"] = str(headline)
    problems, ad_only = [], set()
    for name, text in texts.items():
        missing = []
        for n in sorted(numbers(text)):
            if not (grounded_number(n, own_numbers | ad_numbers) or n in ok or n.rstrip("%") in ok):
                missing.append(n)
        for term, pattern in terms:
            if pattern.search(text) and not pattern.search(own) and term.lower() not in ok:
                if pattern.search(ad):
                    ad_only.add(term)
                else:
                    missing.append(term)
        if missing:
            problems.append(f"{name}: not in your CV, CLAUDE.md, projects.md or the ad: {', '.join(missing)}")
    if problems:
        problems.append("Remove or correct these claims. Only a term the letter doesn't claim as the candidate's "
                        "(a product of the company, a gap named as one) may go into grounding_ok: [...] in "
                        "application.md, and job.md must say why.")
    warnings = [f"from the ad only, not in your files: {', '.join(sorted(ad_only))} "
                "(the letter may name them, never claim them)"] if ad_only else []
    return problems, warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    folder = appdirs.find(a.folder) or Path(a.folder)
    if not folder.is_absolute():
        folder = ROOT / folder
    try:
        meta, bodies = read_application(folder)
        variants = yaml.safe_load(VARIANTS.read_text(encoding="utf-8")) or {}
        layout = load_layout()
        short = meta["short"]
        langs = list(bodies)
        suffix = {lg: f"_{lg.upper()}" if len(langs) > 1 else "" for lg in langs}
        cv_out = folder / f"{FILE_PREFIX}_CV_{short}.docx"
        letter_out = {lg: folder / f"{FILE_PREFIX}_Cover_Letter_{short}{suffix[lg]}.docx" for lg in langs}
        tmp_cv = folder / f"~build_CV_{short}.docx"
        tmp_letter = {lg: folder / f"~build_Letter_{short}{suffix[lg]}.docx" for lg in langs}
        cv_doc, summary = build_cv(meta, variants, layout, tmp_cv)
        letter_docs = {lg: build_letter(meta, bodies[lg], lg, langs, layout, tmp_letter[lg]) for lg in langs}
        renamed = 0
        for d, kind, path in [(cv_doc, "CV", tmp_cv)] + [
                (letter_docs[lg], "Anschreiben" if lg == "de" else "Cover Letter", tmp_letter[lg]) for lg in langs]:
            renamed += rename(d)
            apply_fonts(d)
            set_metadata(d, kind)
            d.save(path)
        if NAME_IN_MASTER and NAME_IN_MASTER != APPLICANT and not renamed:
            raise BuildError(f"name_in_master '{NAME_IN_MASTER}' isn't in the master CV or letter (config.yaml)")
        problems, counts = check(meta, bodies, cv_doc, letter_docs)
        grounding_problems, warnings = grounding(meta, bodies, variants, folder)
        problems += grounding_problems
        if problems:
            raise BuildError("checks failed:\n  - " + "\n  - ".join(problems))
        tmp_cv.replace(cv_out)
        for lg in langs:
            tmp_letter[lg].replace(letter_out[lg])
    except BuildError as e:
        for f in folder.glob("~build_*.docx"):
            f.unlink()
        print(f"FAILED {folder.name}: {e}", file=sys.stderr)
        return 1
    print(f"OK {folder.name}: " + ", ".join(
        f"letter{' ' + lg.upper() if len(langs) > 1 else ''} {n} words" for lg, n in counts.items()))
    for line in summary + warnings:
        print(f"  {line}")
    if not a.no_pdf:
        try:
            engine = pdf_export.export([cv_out, *letter_out.values()], CFG.get("pdf_engine"))
        except pdf_export.PdfError as e:
            print(f"WARNING {folder.name}: no PDFs ({e}); the .docx files are built, page limits unchecked",
                  file=sys.stderr)
            return 0
        if engine == "none":
            print("  no PDFs (pdf_engine: none); page limits unchecked")
            return 0
        if engine == "libreoffice":
            print("  PDFs from LibreOffice: its page count can differ from Word's, open the CV PDF once to check")
        problems = check_pages(cv_out.with_suffix(".pdf"), [p.with_suffix(".pdf") for p in letter_out.values()])
        if problems:
            print(f"FAILED {folder.name}: page check:\n  - " + "\n  - ".join(problems), file=sys.stderr)
            return 1
    return 0


def check_pages(cv_pdf, letter_pdfs):
    """Page limits from the exported PDFs (Word's count is the reliable one; LibreOffice can differ)."""
    import pypdf
    problems = []
    for name, path, limit in [("CV", cv_pdf, CV_MAX_PAGES)] + [(f"letter {p.stem}", p, LETTER_MAX_PAGES)
                                                               for p in letter_pdfs]:
        n = len(pypdf.PdfReader(str(path)).pages)
        if n > limit:
            problems.append(f"{name} has {n} pages (max {limit})")
    if any(p.startswith("CV") for p in problems):
        problems.append("Do NOT drop or shorten a project to fix this without the applicant's approval. Try a shorter "
                        "headline first (about 72 characters, one line); otherwise ask the applicant which project to leave "
                        "out (cv.omit_projects + cv.omit_approved).")
    return problems


if __name__ == "__main__":
    sys.exit(main())
