import docx
import pytest

import cv_map
from cv_map import MapError


def make(tmp_path, lines, bold=()):
    d = docx.Document()
    for text in lines:
        d.add_paragraph().add_run(text).bold = text in bold
    p = tmp_path / "cv.docx"
    d.save(p)
    return docx.Document(str(p))


def test_find_anchors(tmp_path):
    d = make(tmp_path, ["Jane Doe", "", "Headline here", "Berlin, 01.10.2026"])
    paras = cv_map.paragraphs(d)
    assert cv_map.find(paras, {"after": "jane  doe"}, "x").text == "Headline here"   # case and spaces ignored
    assert cv_map.find(paras, {"starts_with": "Berlin,"}, "x").text.startswith("Berlin")
    assert cv_map.find(paras, "Headline here", "x").text == "Headline here"


def test_ambiguous_anchor_asks_for_remap(tmp_path):
    d = make(tmp_path, ["Skills", "Skills"])
    with pytest.raises(MapError, match="remap my CV"):
        cv_map.find(cv_map.paragraphs(d), {"text": "Skills"}, "x")
    with pytest.raises(MapError, match="unknown anchor"):
        cv_map.find(cv_map.paragraphs(d), {"near": "Skills"}, "x")


def test_section_rules(tmp_path):
    d = make(tmp_path, ["PROJECTS", "Alpha, a tool", "about alpha", "", "Beta, another", "about beta", "",
                        "SKILLS", "Tools: Excel", "Languages: German", "", "EDUCATION", "B.A."],
             bold=["Alpha, a tool", "Beta, another"])
    names, blocks = cv_map.section(d, "projects", {"start": "PROJECTS", "end": ["SKILLS"], "items": "blank_line",
                                                   "name_split": ","})
    assert names == ["Alpha", "Beta"]
    assert len(blocks["Alpha"]) == 3 and len(blocks["Beta"]) == 3     # every item ends with its blank line
    names, _ = cv_map.section(d, "projects", {"start": "PROJECTS", "end": ["SKILLS"], "items": "bold",
                                              "name_split": ","})
    assert names == ["Alpha", "Beta"]
    names, blocks = cv_map.section(d, "skills", {"start": "SKILLS", "end": ["EDUCATION"], "items": "colon",
                                                 "name_split": ":"})
    assert names == ["Tools", "Languages"]
    assert len(blocks["Languages"]) == 1                              # trailing spacer stays put


def test_section_errors(tmp_path):
    d = make(tmp_path, ["SKILLS", "Tools: Excel", "Tools: Word"])
    with pytest.raises(MapError, match="two items"):
        cv_map.section(d, "skills", {"start": "SKILLS", "items": "colon", "name_split": ":"})
    with pytest.raises(MapError, match="items must be"):
        cv_map.section(d, "skills", {"start": "SKILLS", "items": "commas"})


def test_letter_parts(tmp_path):
    d = make(tmp_path, ["Jane Doe", "Berlin, 01.10.2026", "Acme GmbH", "Application: X", "Dear team,",
                        "First.", "Second.", "", "Kind regards,"])
    parts = cv_map.letter_parts(d, {"date": {"starts_with": "Berlin,"}, "recipient": {"starts_with": "Acme"},
                                    "subject": {"starts_with": "Application:"}, "greeting": {"starts_with": "Dear"}})
    assert [p.text for p in parts["body"]] == ["First.", "Second."]
    with pytest.raises(MapError, match="greeting"):
        cv_map.letter_parts(d, {"date": {"starts_with": "Berlin,"}})


def test_load_rejects_unknown_keys(tmp_path):
    (tmp_path / "layout.yaml").write_text("cv: {}\nresume: {}\n", encoding="utf-8")
    with pytest.raises(MapError, match="resume"):
        cv_map.load(tmp_path / "layout.yaml")
    with pytest.raises(MapError, match="not found"):
        cv_map.load(tmp_path / "missing.yaml")


def test_not_list_rule_and_name_split_list(tmp_path):
    d = docx.Document()
    for text, bullet in (("PROJEKTE", False), ("Alpha, Thesis\t2025", False), ("did a", True), ("did b", True),
                         ("Beta\t2024", False), ("did c", True)):
        d.add_paragraph(text, style="List Bullet" if bullet else None)
    p = tmp_path / "cv.docx"
    d.save(p)
    names, blocks = cv_map.section(docx.Document(str(p)), "projects",
                                   {"start": "PROJEKTE", "items": "not_list", "name_split": ["\t", ","]})
    assert names == ["Alpha", "Beta"] and len(blocks["Alpha"]) == 3


def test_letter_closing_and_address_block(tmp_path):
    d = make(tmp_path, ["Acme GmbH", "Street 1", "12345 Town", "", "Berlin, 01.10.2026", "Dear team,",
                        "First.", "Second.", "Kind regards,", "Jane"])
    lmap = {"recipient": {"text": "Acme GmbH"}, "date": {"starts_with": "Berlin,"}, "greeting": {"starts_with": "Dear"},
            "closing": {"starts_with": "Kind regards"}}
    parts = cv_map.letter_parts(d, lmap)
    assert [p.text for p in parts["recipient"]] == ["Acme GmbH", "Street 1", "12345 Town"]
    assert [p.text for p in parts["body"]] == ["First.", "Second."]
    with pytest.raises(MapError, match="unknown part"):
        cv_map.letter_parts(d, {**lmap, "signature": {"text": "Jane"}})
