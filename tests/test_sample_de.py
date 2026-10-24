"""End to end on the German sample workspace: plain-paragraph CV, multi-line address, closing without an empty line."""
import json

import docx

APP = "applications/2026-10-23_Beispiel-Werke_Berechnungsingenieur"


def texts(path):
    return [p.text for p in docx.Document(str(path)).paragraphs]


def rewrite_letter(root, recipient=None, body=None):
    """Rewrite the master letter with another address block or body (lists of paragraphs)."""
    from conftest import SAMPLE_DE
    import importlib.util
    spec = importlib.util.spec_from_file_location("make_docs_de", SAMPLE_DE / "make_docs.py")
    md = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(md)
    lines = list(md.LETTER)
    if recipient is not None:
        lines[4:7] = recipient
    if body is not None:
        g = next(i for i, t in enumerate(lines) if t.startswith("Sehr geehrte"))
        c = lines.index("Mit freundlichen Grüßen")
        lines[g + 1:c] = body
    d = docx.Document()
    for text in lines:
        d.add_paragraph(text)
    d.save(root / "cv" / "Jonas_Becker_Anschreiben.docx")


def test_setup_check_ready(tool_de):
    result = json.loads(tool_de("setup_check.py", "--init", "--json").stdout)
    assert result["ready"], [c for c in result["checks"] if c["status"] == "todo"]
    assert any(c["item"] == "CV map matches the master CV and letter" and c["status"] == "ok"
               for c in result["checks"])


def test_cv_map_check(tool_de):
    r = tool_de("cv_map.py", "check")
    assert r.returncode == 0, r.stdout
    assert "date: Regensburg, 01.10.2026" in r.stdout
    assert ("section projects: ['FEM-Analyse Radträger', 'Formula Student Fahrwerk', "
            "'Python-Skript für Kinematikauswertung']") in r.stdout
    assert "letter recipient: 3 line(s): Beispiel Werke AG / Hauptstraße 1 / 80331 München" in r.stdout
    assert "letter closing: Mit freundlichen Grüßen" in r.stdout
    assert "letter body: 2 paragraphs" in r.stdout


def test_build(tool_de, sample_de):
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 0, r.stderr
    folder = sample_de / APP
    cv = texts(folder / "Jonas_Becker_CV_Beispiel-Werke.docx")
    assert cv[1] == "Junior Berechnungsingenieur | FEM, ANSYS & Abaqus"
    assert cv[-1] == "Regensburg, 23.10.2026"                         # the date line follows the letter date
    titles = [t.split("\t")[0] for t in cv[cv.index("PROJEKTE") + 1:cv.index("AUSBILDUNG")] if "\t" in t]
    assert titles == ["Python-Skript für Kinematikauswertung", "FEM-Analyse Radträger, Bachelorarbeit",
                      "Formula Student Fahrwerk"]
    fem = next(i for i, t in enumerate(cv) if t.startswith("FEM-Analyse"))
    assert cv[fem + 1].startswith("Topologieoptimierung") and cv[fem + 2].startswith("Validierung")
    tech = [t.split(":")[0] for t in cv[cv.index("TECHNISCHE KOMPETENZEN") + 1:cv.index("BERUFSERFAHRUNG")] if ":" in t]
    assert tech == ["Simulation", "Programmierung", "CAD", "Sprachen"]

    assert "Jonas Becker" in cv[0] and not any("Michael" in t for t in cv)    # name_in_master
    letter = texts(folder / "Jonas_Becker_Cover_Letter_Beispiel-Werke_DE.docx")
    start = letter.index("Beispiel Werke AG")
    assert letter[start:start + 6] == ["Beispiel Werke AG", "Personalabteilung", "Frau Muster", "Werkstraße 12",
                                       "90402 Nürnberg", ""]                  # 5 lines in place of 3, nothing left over
    assert "Hauptstraße 1" not in letter and "80331 München" not in letter
    assert "Regensburg, 23.10.2026" in letter
    greet = letter.index("Sehr geehrte Frau Muster,")
    close = letter.index("Mit freundlichen Grüßen")
    assert close - greet - 1 == 5                                             # 5 body paragraphs, closing kept
    assert letter[close + 1] == "Jonas Becker" and letter[0] == "Jonas Becker"
    assert not any("ersetzt" in t for t in letter)

    en = texts(folder / "Jonas_Becker_Cover_Letter_Beispiel-Werke_EN.docx")
    assert "Application: Junior Structural Analysis Engineer (m/f/d)" in en and "Dear Ms Muster," in en
    close = en.index("Kind regards")                                          # the closing follows the language
    assert close - en.index("Dear Ms Muster,") - 1 == 5 and en[close + 1] == "Jonas Becker"
    assert not (folder / "Jonas_Becker_Cover_Letter_Beispiel-Werke.docx").exists()


def test_languages_need_every_body(tool_de, sample_de):
    app = sample_de / APP / "application.md"
    app.write_text(app.read_text(encoding="utf-8").split("<!-- en -->")[0], encoding="utf-8")
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 1 and "no letter body for 'en'" in r.stderr


def test_one_language_and_student_checks_in_german(tool_de, sample_de):
    app = sample_de / APP / "application.md"
    text = app.read_text(encoding="utf-8").replace("languages: [de, en]", "languages: [de]")
    app.write_text(text.replace("role_type: Junior", "role_type: Werkstudent").replace("ab sofort", "bald"),
                   encoding="utf-8")
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 1 and "'ab sofort'" in r.stderr and "stunden pro woche" in r.stderr
    assert "letter EN" not in r.stderr


def test_name_in_master_must_exist(tool_de, sample_de):
    cfg = sample_de / "config.yaml"
    cfg.write_text(cfg.read_text(encoding="utf-8").replace("Jonas Michael Becker", "Jonas M. Becker"), encoding="utf-8")
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 1 and "name_in_master 'Jonas M. Becker' isn't in the master" in r.stderr


def test_recipient_in_one_paragraph(tool_de, sample_de):
    """An address written with line breaks inside one paragraph gets the new lines the same way."""
    rewrite_letter(sample_de, ["Beispiel Werke AG\nHauptstraße 1\n80331 München"])
    layout = sample_de / "cv" / "layout.yaml"
    layout.write_text(layout.read_text(encoding="utf-8").replace('{text: "Beispiel Werke AG"}',
                                                                 '{starts_with: "Beispiel Werke AG"}'), encoding="utf-8")
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 0, r.stderr
    letter = texts(sample_de / APP / "Jonas_Becker_Cover_Letter_Beispiel-Werke_DE.docx")
    assert "Beispiel Werke AG\nPersonalabteilung\nFrau Muster\nWerkstraße 12\n90402 Nürnberg" in letter


def test_spacer_between_body_paragraphs(tool_de, sample_de):
    """Empty lines between the master's body paragraphs are kept between the new ones."""
    rewrite_letter(sample_de, body=["Erster Absatz.", "", "Zweiter Absatz.", ""])
    r = tool_de("build_application.py", APP, "--no-pdf")
    assert r.returncode == 0, r.stderr
    letter = texts(sample_de / APP / "Jonas_Becker_Cover_Letter_Beispiel-Werke_DE.docx")
    body = letter[letter.index("Sehr geehrte Frau Muster,") + 1:letter.index("Mit freundlichen Grüßen")]
    assert len(body) == 10 and body[1] == "" and body[-1] == ""              # 5 paragraphs, each followed by a gap
